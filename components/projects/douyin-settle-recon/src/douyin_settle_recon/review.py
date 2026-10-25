"""Read-only workbook comparison and a loopback-only review server."""
from __future__ import annotations

from collections import Counter, defaultdict, deque
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
import csv
import io
import json
from pathlib import Path
import re
import subprocess
import sys
import threading
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

NAME = re.compile(r"^(\d{4}-?(?:0[1-9]|1[0-2]))【抖音】(.+)\.xlsx$")
GROUPS = [
    ("fees", "费用分类汇总", "核对汇总", "summary"),
    ("checks", "结算金额校验", "核对汇总", "summary"),
    ("balances", "账户余额", "核对汇总", "summary"),
    ("identity", "店铺编码", "核对汇总", "summary"),
    ("processed", "资金账单（核对）", "资金账单 (核对)", "table"),
    ("funds", "原始资金账单", "资金账单", "table"),
    ("settlement", "结算账单", "结算账单", "table"),
    ("deposit", "保证金账户", "保证金账户", "table"),
    ("evidence", "余额与管家证据", "", "evidence"),
    ("template", "模板结构与公式", "", "template"),
]
ALIASES = {"店铺劵": "店铺券", "动帐流水号": "动账流水号"}


def scalar(value):
    if isinstance(value, (date, datetime)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, Decimal):
        return str(value)
    return value


def text(value):
    return "" if value is None else str(scalar(value)).strip().lstrip("'").strip()


def number(value):
    if value is None or isinstance(value, bool):
        return None
    try:
        n = Decimal(str(value))
        return n if n.is_finite() else None
    except (InvalidOperation, ValueError):
        return None


def cents(value):
    n = number(value)
    return n.quantize(Decimal(".01"), rounding=ROUND_HALF_UP) if n is not None else None


def equivalent(a, b, field=""):
    if a is None or b is None:
        return a is b
    # Identifiers never travel through floating point or monetary rounding.
    if any(k in field for k in ("单号", "流水号", "编号", "编码", "ID")):
        return text(a) == text(b)
    if isinstance(a, (int, float, Decimal)) or isinstance(b, (int, float, Decimal)):
        na, nb = cents(a), cents(b)
        if na is not None and nb is not None:
            return na == nb
    return text(a) == text(b)


def item(label, a, b, expected_location="", actual_location="", field="", reason="", account="", forced=False):
    same = equivalent(a, b, field or label) and not forced
    na, nb = cents(a), cents(b)
    delta = str(nb - na) if na is not None and nb is not None else None
    if not same and not reason:
        if a is None or b is None:
            reason = "预期为空，实际有值" if a is None else "实际为空（空值不按 0 处理）"
        elif str(b).startswith("#"):
            reason = "实际结果包含 Excel 错误值"
        else:
            reason = "金额不一致" if delta is not None else "内容不一致"
    return {"label": label, "field": field or label, "account": account,
            "expected": scalar(a), "actual": scalar(b), "delta": delta,
            "expected_location": expected_location, "actual_location": actual_location,
            "status": "same" if same else "diff", "reason": reason}


def finish(items, **meta):
    different = sum(r["status"] == "diff" for r in items)
    return {"items": items, "total": len(items), "different": different,
            "same": len(items) - different, "status": "diff" if different else "same", **meta}


def rows(ws):
    return [(n, tuple(r)) for n, r in enumerate(ws.iter_rows(values_only=True), 1)
            if any(v is not None for v in r)]


def summary_compare(expected, actual, group):
    items = []
    if "核对汇总" not in expected.sheetnames or "核对汇总" not in actual.sheetnames:
        return finish([item("核对汇总工作表", "核对汇总" if "核对汇总" in expected.sheetnames else None,
                            "核对汇总" if "核对汇总" in actual.sheetnames else None,
                            reason="缺少核对汇总，无法完成此项核对", forced=True)])
    e = dict(rows(expected["核对汇总"]))
    a = dict(rows(actual["核对汇总"]))
    get = lambda grid, r, c: grid.get(r, ())[c-1] if len(grid.get(r, ())) >= c else None
    if group == "fees":
        def fees(grid):
            out = {}
            for r in grid:
                account, category = get(grid,r,2), get(grid,r,3)
                if r > 1 and category is not None:
                    key = (text(account), text(category))
                    if key in out:
                        raise ValueError(f"核对汇总存在重复费用项目：{key}")
                    out[key] = (r, get(grid,r,4))
            return out
        ef, af = fees(e), fees(a)
        for key in dict.fromkeys([*ef, *af]):
            er, ev = ef.get(key, (None,None)); ar, av = af.get(key, (None,None))
            items.append(item(key[1], ev, av, f"核对汇总!D{er}" if er else "",
                              f"核对汇总!D{ar}" if ar else "", account=key[0],
                              reason="缺少费用项目" if ar is None else "新增费用项目" if er is None else "",
                              forced=er is None or ar is None))
        def total(values):
            amounts = [cents(v[1]) for v in values.values()]
            return str(sum(amounts,Decimal(0))) if amounts and all(v is not None for v in amounts) else None
        totals = {"expected_total": total(ef), "actual_total": total(af)}
        if not ef or not af:
            items.append(item("费用分类数据", "有费用项目" if ef else None, "有费用项目" if af else None,
                              reason="费用汇总为空，无法确认结果",forced=True))
        return finish(items, **totals)
    if group == "checks":
        for r in range(2,5):
            items.append(item(text(get(e,r,7)) or f"校验第 {r} 行", get(e,r,8), get(a,r,8),
                              f"核对汇总!H{r}", f"核对汇总!H{r}"))
    elif group == "identity":
        # Compare every populated fee row: a code missing from even one row is visible.
        for r in sorted(set(e) | set(a)):
            if r > 1 and (get(e,r,3) is not None or get(a,r,3) is not None):
                items.append(item(text(get(e,r,3)) or text(get(a,r,3)), get(e,r,1),get(a,r,1),
                                  f"核对汇总!A{r}",f"核对汇总!A{r}", field="店铺编码",account=text(get(e,r,2))))
    else:
        def accounts(grid):
            return {text(get(grid,r,12)):r for r in grid if r > 1 and get(grid,r,12) is not None}
        ea, aa = accounts(e), accounts(a)
        for account in dict.fromkeys([*ea,*aa]):
            er, ar = ea.get(account), aa.get(account)
            for c,label in [(13,"期初余额"),(14,"本期发生额"),(15,"期末余额"),(16,"余额校验")]:
                items.append(item(label,get(e,er,c),get(a,ar,c),
                                  f"核对汇总!{get_column_letter(c)}{er}" if er else "",
                                  f"核对汇总!{get_column_letter(c)}{ar}" if ar else "",account=account,
                                  forced=er is None or ar is None))
    return finish(items)


def canonical(header):
    h = text(header)
    return ALIASES.get(h, h)


def read_table(ws):
    data = iter(ws.iter_rows(values_only=True))
    first = next(data, ())
    columns = {}
    counts = Counter()
    for i, header in enumerate(first):
        if header is not None:
            name = text(header); counts[name] += 1
            columns[name if counts[name] == 1 else f"{name}（{counts[name]}）"] = i
    records = [(rn, row) for rn, row in enumerate(data,2) if any(v is not None for v in row)]
    named_positions = set(columns.values())
    for _, row in records:
        for i, value in enumerate(row):
            if value is not None and i not in named_positions:
                columns[f"未命名列（{get_column_letter(i+1)}）"] = i
                named_positions.add(i)
    return columns, records


def align_aliases(expected_columns, actual_columns):
    """Keep exact names when aliases coexist; never choose an empty alias over an exact field."""
    def counts(columns):
        return Counter(canonical(c) for c in columns)
    ec, ac = counts(expected_columns), counts(actual_columns)
    def rename(columns):
        return {(canonical(c) if ec[canonical(c)] <= 1 and ac[canonical(c)] <= 1 else c):i
                for c,i in columns.items()}
    return rename(expected_columns), rename(actual_columns)


def table_compare(expected, actual, sheet):
    if sheet not in expected.sheetnames or sheet not in actual.sheetnames:
        return finish([item(sheet, sheet if sheet in expected.sheetnames else None,
                            sheet if sheet in actual.sheetnames else None, forced=True,
                            reason="工作表缺失，无法核对明细")], expected_rows=0,actual_rows=0,checked_cells=0)
    ec, erows = read_table(expected[sheet]); ac, arows = read_table(actual[sheet])
    ec, ac = align_aliases(ec, ac)
    items = []
    for field in dict.fromkeys([*ec,*ac]):
        if field not in ec or field not in ac:
            items.append(item("表头",field if field in ec else None,field if field in ac else None,
                              f"{sheet}!{get_column_letter(ec[field]+1)}1" if field in ec else "",
                              f"{sheet}!{get_column_letter(ac[field]+1)}1" if field in ac else "",
                              field=field,reason="实际缺少字段" if field not in ac else "实际新增字段",forced=True))
    common = [c for c in ec if c in ac]
    if not common:
        return finish(items or [item(sheet,"有表头","无可匹配表头",forced=True)],
                      expected_rows=len(erows),actual_rows=len(arows),checked_cells=0)
    candidates = [("动账流水号",), ("操作单号",), ("子订单号","订单号","结算时间"), ("订单号","结算时间")]
    key_fields = next((ks for ks in candidates if all(k in common for k in ks)), ())
    value = lambda row, col: row[col] if col < len(row) else None
    def group_records(records, cols):
        out = defaultdict(list)
        for rn, row in records:
            key = tuple(text(value(row, cols[k])) for k in key_fields) if key_fields else (str(rn),)
            if not any(key):
                key = (f"空业务编号·第{rn}行",)
            out[key].append((rn,row))
        return out
    eg, ag = group_records(erows, ec), group_records(arows, ac)
    def signature(row, cols):
        result = []
        for field in common:
            v = value(row, cols[field])
            if v is None: result.append(None)
            elif isinstance(v,(int,float,Decimal)) and not any(k in field for k in ("单号","流水号","编号","编码","ID")):
                result.append(str(cents(v)))
            else: result.append(text(v))
        return tuple(result)
    checked, matched_rows = 0, 0
    for key in dict.fromkeys([*eg,*ag]):
        # Cancel exact duplicate records first; remaining duplicates pair in source order.
        left, right = eg.get(key,[]), ag.get(key,[])
        indexed = defaultdict(deque)
        for rn,row in right: indexed[signature(row,ac)].append((rn,row))
        pending = []
        for rn,row in left:
            bucket = indexed.get(signature(row,ec))
            if bucket:
                bucket.popleft(); checked += len(common); matched_rows += 1
            else: pending.append((rn,row))
        remaining = sorted((r for bucket in indexed.values() for r in bucket), key=lambda r:r[0])
        for i in range(max(len(pending),len(remaining))):
            er, ev = pending[i] if i < len(pending) else (None,())
            ar, av = remaining[i] if i < len(remaining) else (None,())
            label = " / ".join(key)
            if er is None or ar is None:
                excerpt = lambda row,cols: json.dumps({k:scalar(value(row,cols[k])) for k in cols if value(row,cols[k]) is not None},ensure_ascii=False)
                items.append(item(label,excerpt(ev,ec) if er else None,excerpt(av,ac) if ar else None,
                                  f"{sheet}!A{er}" if er else "",f"{sheet}!A{ar}" if ar else "",
                                  field="整行",reason="实际缺少此行" if ar is None else "实际新增此行",forced=True))
                continue
            for field in common:
                checked += 1
                e, a = value(ev,ec[field]), value(av,ac[field])
                if not equivalent(e,a,field):
                    items.append(item(label,e,a,f"{sheet}!{get_column_letter(ec[field]+1)}{er}",
                                      f"{sheet}!{get_column_letter(ac[field]+1)}{ar}",field=field))
    return finish(items, expected_rows=len(erows),actual_rows=len(arows),checked_cells=checked,
                  matched_rows=matched_rows, key_fields=list(key_fields),
                  duplicate_keys=sum(len(v)>1 for v in eg.values()), differences_only=True,
                  note="按业务编号匹配，忽略行列顺序；重复编号先消除完全相同的行，其余按原顺序逐项比较。业务编号为空时按原始行号定位。只存储差异；空行忽略。")


def evidence_compare(expected,actual):
    items=[]
    for sheet in ("余额账单","管家账户"):
        aliases=("余额截图","余额账单") if sheet=="余额账单" else (sheet,)
        expected_sheet=next((n for n in aliases if n in expected.sheetnames),sheet)
        actual_sheet=next((n for n in aliases if n in actual.sheetnames),sheet)
        source = {f"{get_column_letter(c)}{rn}":v for rn,row in rows(expected[expected_sheet])
                  for c,v in enumerate(row,1) if v is not None} if expected_sheet in expected.sheetnames else {}
        target = {}
        if actual_sheet in actual.sheetnames:
            target = {f"{get_column_letter(c)}{rn}":v for rn,row in rows(actual[actual_sheet])
                      for c,v in enumerate(row,1) if v is not None}
        elif sheet+"原始记录" in actual.sheetnames:
            target = {text(row[0]):row[1] for rn,row in rows(actual[sheet+"原始记录"]) if rn>1 and len(row)>1}
        for coord in dict.fromkeys([*source,*target]):
            items.append(item(sheet,source.get(coord),target.get(coord),f"{expected_sheet}!{coord}",
                              f"{actual_sheet}!{coord}" if actual_sheet in actual.sheetnames else f"{sheet}原始记录 · {coord}"))
    return finish(items,note="仅核对工作表内的单元格数值。嵌入截图的图像内容、历史时点和余额有效性仍需打开原表人工确认。")


def formula_text(value):
    return value if isinstance(value,str) else getattr(value,"text",None)


def template_compare(template_path, actual_path):
    items=[]
    with_workbooks=[]
    try:
        for path in (template_path,actual_path):
            with_workbooks.append(load_workbook(path,read_only=True,data_only=False))
        template, actual = with_workbooks
        items.append(item("工作表名称与顺序",json.dumps(template.sheetnames,ensure_ascii=False),
                          json.dumps(actual.sheetnames,ensure_ascii=False),reason="工作表名称或顺序与模板不同" if template.sheetnames!=actual.sheetnames else ""))
        for sheet in ("核对汇总","资金账单 (核对)"):
            if sheet not in actual.sheetnames:
                items.append(item(sheet,sheet,None,reason="缺少模板要求的工作表",forced=True));continue
            ts, ass = template[sheet], actual[sheet]
            trows=list(ts.iter_rows(max_row=43 if sheet=="核对汇总" else 2))
            arows=list(ass.iter_rows(max_row=43 if sheet=="核对汇总" else 2))
            for rn,row in enumerate(trows,1):
                for col,cell in enumerate(row,1):
                    if rn==1 and cell.value is not None:
                        av=arows[rn-1][col-1].value if rn<=len(arows) and col<=len(arows[rn-1]) else None
                        items.append(item(f"{sheet} · 表头",cell.value,av,f"{sheet}!{cell.coordinate}",f"{sheet}!{cell.coordinate}"))
                    elif cell.data_type=="f" and (sheet=="核对汇总" or col in (40,41,42)):
                        av=arows[rn-1][col-1].value if rn<=len(arows) and col<=len(arows[rn-1]) else None
                        items.append(item(f"{sheet} · 模板公式",formula_text(cell.value),formula_text(av),
                                          f"{sheet}!{cell.coordinate}",f"{sheet}!{cell.coordinate}",field="公式"))
        return finish(items,note="以处理模版核对工作表名称与顺序、固定表头、汇总公式和资金加工第 2 行的公式。后续行的公式填充、格式及截图布局不在此项检查范围。")
    finally:
        for wb in with_workbooks:wb.close()


class ReviewStore:
    def __init__(self,root,run_roots=None,expected_dir=None,template=None):
        self.root=Path(root).resolve()
        self.run_roots=[Path(p).resolve() for p in (run_roots or [self.root/"runs",self.root/"data/runs"])]
        self.expected_dir=Path(expected_dir or self.root/"doc").resolve()
        self.template=Path(template or self.root/"doc/【抖音】月度账单处理模版.xlsx").resolve()
        self.pool=ThreadPoolExecutor(max_workers=1)
        self.lock=threading.Lock();self.jobs={};self.entries={}

    def catalog(self):
        candidates=defaultdict(list)
        for root in self.run_roots:
            if not root.is_dir():continue
            for p in root.rglob("*.xlsx"):
                relative=p.relative_to(root)
                if any(part in ("raw","rules",".review") for part in relative.parts[:-1]):continue
                if NAME.fullmatch(p.name) and not p.is_symlink():candidates[p.name].append(p.resolve())
        result=[]
        for name,paths in candidates.items():
            # A run's timestamp wins over a later touch/copy of an older workbook.
            def rank(p):
                runstamp=re.search(r"\d{8}T\d{6}Z",str(p.parent))
                return (datetime.strptime(runstamp[0],"%Y%m%dT%H%M%SZ").replace(tzinfo=timezone.utc).timestamp() if runstamp else p.stat().st_mtime,p.stat().st_mtime)
            paths.sort(key=rank,reverse=True);actual=paths[0]
            match=NAME.fullmatch(name)
            stem=Path(name).stem
            expected_candidates=[self.expected_dir/f"{stem}-预期结果.xlsx",
                                 self.expected_dir/f"{stem}预期结果.xlsx",self.expected_dir/name]
            expected=next((p for p in expected_candidates if p.is_file()),expected_candidates[0])
            ident=name
            month=match[1].replace('-', '')
            shop=match[2][2:] if match[2].startswith('Y-') else match[2]
            entry={"id":ident,"name":name,"month":f"{month[:4]}-{month[4:]}","shop":shop,
                   "actual_path":str(actual),"expected_path":str(expected) if expected.is_file() else None,
                   "run_id":actual.parent.name,"modified":datetime.fromtimestamp(actual.stat().st_mtime).isoformat(timespec="seconds"),
                   "versions":len(paths),"has_expected":expected.is_file()}
            self.entries[ident]=entry;result.append(entry)
        result.sort(key=lambda e:e["modified"],reverse=True)
        self.entries={e["id"]:e for e in result}
        return {"entries":result,"run_roots":[str(p) for p in self.run_roots],"template":str(self.template),
                "groups":[{"id":g,"label":label,"sheet":sheet,"kind":kind} for g,label,sheet,kind in GROUPS]}

    def entry(self,ident):
        entry=self.entries.get(ident)
        if not entry:raise ValueError("未找到此账单，请重新扫描目录")
        return entry

    def request(self,ident,group):
        entry=self.entry(ident)
        if group not in {g[0] for g in GROUPS}:raise ValueError("未知核对项目")
        if not entry["has_expected"]:raise ValueError("doc 中没有匹配的预期表（支持文件名末尾添加 -预期结果），不能判定核对通过")
        ep,ap=Path(entry["expected_path"]),Path(entry["actual_path"])
        stamp=lambda p:(str(p),p.stat().st_size,p.stat().st_mtime_ns)
        key=(ident,group,stamp(ep),stamp(ap),stamp(self.template) if group=="template" else None)
        with self.lock:
            if key not in self.jobs:
                self.jobs[key]=self.pool.submit(self.compare,entry.copy(),group)
            job=self.jobs[key]
        if not job.done():return {"state":"loading","group":group}
        try:return {"state":"ready","group":group,**job.result()}
        except Exception as exc:return {"state":"error","error":f"核对失败：{exc}"}

    def compare(self,entry,group):
        ep,ap=Path(entry["expected_path"]),Path(entry["actual_path"])
        if group=="template":return template_compare(self.template,ap)
        before=[(p.stat().st_size,p.stat().st_mtime_ns) for p in (ep,ap)]
        expected=load_workbook(ep,read_only=True,data_only=True)
        actual=None
        try:
            actual=load_workbook(ap,read_only=True,data_only=True)
            kind=next(g for g in GROUPS if g[0]==group)
            if kind[3]=="summary":result=summary_compare(expected,actual,group)
            elif kind[3]=="evidence":result=evidence_compare(expected,actual)
            else:result=table_compare(expected,actual,kind[2])
            result["sheets"]={"expected":expected.sheetnames,"actual":actual.sheetnames}
            result["compared_at"]=datetime.now().isoformat(timespec="seconds")
            if before != [(p.stat().st_size,p.stat().st_mtime_ns) for p in (ep,ap)]:
                raise ValueError("文件在核对过程中发生变化，请重新核对")
            return result
        finally:
            expected.close()
            if actual is not None:actual.close()


def page_result(result,params):
    if result.get("state")!="ready":return result
    result=result.copy();items=result.pop("items")
    query=params.get("q",[""])[0].strip().casefold();status=params.get("status",["all"])[0]
    filtered=[r for r in items if (status=="all" or r["status"]==status)
              and (not query or query in json.dumps(r,ensure_ascii=False).casefold())]
    offset=max(0,int(params.get("offset",[0])[0]));limit=min(200,max(1,int(params.get("limit",[100])[0])))
    return {**result,"items":filtered[offset:offset+limit],"filtered":len(filtered),"offset":offset,"limit":limit}


def csv_data(result):
    stream=io.StringIO();writer=csv.writer(stream)
    writer.writerow(["项目 / 业务编号","账户","字段","预期值","实际值","差额（实际－预期）","结论","预期位置","实际位置","原因"])
    for r in result["items"]:
        if r["status"]!="diff":continue
        def safe(v):
            s="" if v is None else str(v)
            return "'"+s if s.lstrip().startswith(("=","+","-","@")) else s
        writer.writerow([safe(r[k]) for k in ("label","account","field","expected","actual","delta","status","expected_location","actual_location","reason")])
    return ("\ufeff"+stream.getvalue()).encode("utf-8")


def serve(root,host="127.0.0.1",port=5189,run_roots=None,expected_dir=None,template=None):
    if host not in ("127.0.0.1","localhost","::1"):
        raise ValueError("核对页面仅支持本机访问")
    store=ReviewStore(root,run_roots,expected_dir,template);store.catalog()
    web=Path(__file__).parent/"web"
    class Handler(BaseHTTPRequestHandler):
        def log_message(self,fmt,*args):pass
        def send(self,status,body,kind="application/json; charset=utf-8",headers=None):
            if not isinstance(body,bytes):body=json.dumps(body,ensure_ascii=False,default=scalar).encode()
            self.send_response(status);self.send_header("Content-Type",kind)
            self.send_header("Content-Length",str(len(body)));self.send_header("Cache-Control","no-store")
            self.send_header("X-Content-Type-Options","nosniff")
            for k,v in (headers or {}).items():self.send_header(k,v)
            self.end_headers();self.wfile.write(body)
        def local(self):
            host_header=self.headers.get("Host","")
            allowed={f"127.0.0.1:{self.server.server_port}",f"localhost:{self.server.server_port}"}
            origin=self.headers.get("Origin")
            return host_header in allowed and (not origin or origin in {"http://"+h for h in allowed})
        def do_GET(self):
            if not self.local():self.send(403,{"error":"仅允许本机同源访问"});return
            parsed=urlparse(self.path);params=parse_qs(parsed.query)
            try:
                if parsed.path=="/api/catalog":self.send(200,store.catalog())
                elif parsed.path in ("/api/compare","/api/export"):
                    result=store.request(params.get("id",[""])[0],params.get("group",["fees"])[0])
                    if parsed.path=="/api/export" and result.get("state")=="ready":
                        self.send(200,csv_data(result),"text/csv; charset=utf-8",{"Content-Disposition":"attachment; filename=review-differences.csv"})
                    else:self.send(200,page_result(result,params))
                elif parsed.path=="/api/issues":
                    entry=store.entry(params.get("id",[""])[0]);path=Path(entry["actual_path"]).parent/"result.json"
                    data=json.loads(path.read_text()) if path.is_file() else {}
                    self.send(200,{"context":data.get("context",{}),"checks":data.get("checks",[]),"issues":data.get("issues",[]),"available":path.is_file()})
                elif parsed.path in ("/","/app.js","/style.css"):
                    file=web/({"/":"index.html"}.get(parsed.path,parsed.path.lstrip("/")))
                    self.send(200,file.read_bytes(),{".html":"text/html; charset=utf-8",".js":"text/javascript; charset=utf-8",".css":"text/css; charset=utf-8"}[file.suffix])
                else:self.send(404,{"error":"页面不存在"})
            except (ValueError,OSError,KeyError) as exc:self.send(400,{"error":str(exc)})
        def do_POST(self):
            if not self.local():self.send(403,{"error":"仅允许本机同源访问"});return
            try:
                if self.path!="/api/open":self.send(404,{"error":"未知操作"});return
                size=int(self.headers.get("Content-Length","0"))
                if not 0<size<8192:raise ValueError("无效请求")
                if self.headers.get_content_type()!="application/json":raise ValueError("需要 JSON 请求")
                body=json.loads(self.rfile.read(size));entry=store.entry(body.get("id",""))
                side=body.get("side")
                if side not in ("expected","actual"):raise ValueError("未知表格")
                path=entry.get(side+"_path")
                if not path or not Path(path).is_file():raise ValueError("表格文件不存在")
                command=["open",path] if sys.platform=="darwin" else ["cmd","/c","start","",path] if sys.platform=="win32" else ["xdg-open",path]
                subprocess.run(command,check=True,timeout=15,capture_output=True)
                self.send(200,{"ok":True,"path":path})
            except (ValueError,OSError,subprocess.SubprocessError) as exc:self.send(400,{"error":f"打开失败：{exc}"})
    server=ThreadingHTTPServer((host,port),Handler)
    print(f"抖音账单核对页面：http://{host}:{server.server_port}",flush=True)
    try:server.serve_forever()
    except KeyboardInterrupt:pass
    finally:server.server_close();store.pool.shutdown(wait=False,cancel_futures=True)
