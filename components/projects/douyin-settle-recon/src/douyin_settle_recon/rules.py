from __future__ import annotations

from collections import defaultdict
from decimal import Decimal
from pathlib import Path
import json
import re

import openpyxl

from .model import Column, DataError, Dataset, Row, Table, decimal, sha256
from .parse import target_rows
from .schema import ACCOUNTS, ALIASES, FUND_COLUMNS, DATE_FIELDS, AMOUNT_FIELDS
from .qianchuan import TOTAL_FORMULAS, prepare as prepare_qianchuan

AP_BRANCHES = [('运费','运费赔付'),('管家账户充值','管家账户充值'),
               ('聚合账户充值至管家账户','管家账户充值'),('充值保证金','保证金充值'),
               ('充值巨量千川','巨量千川充值'),('超','超时赔付')]


class TemplateRules:
    def __init__(self, path: Path, overrides: Path | None = None):
        self.path = path
        self.template_sha256 = sha256(path)
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            s = wb['资金账单 (核对)']
            names = [c.value for c in next(s.iter_rows(max_row=1))][:39]
            if names != FUND_COLUMNS:
                raise DataError('template input columns changed; review before calculating')
            self.formulas = {pos: getattr(s[pos].value, 'text', s[pos].value) for pos in ['AN2','AO2','AP2']}
            expected_an = '=_xlfn.IFS(C2="入账",D2,C2="出账",-D2,TRUE,D2)'
            expected_ap = ('=_xlfn.IFS(ISNUMBER(SEARCH("运费",AM2)),"运费赔付",'
                'ISNUMBER(SEARCH("管家账户充值",AM2)),"管家账户充值",'
                'ISNUMBER(SEARCH("聚合账户充值至管家账户",AM2)),"管家账户充值",'
                'ISNUMBER(SEARCH("充值保证金",AM2)),"保证金充值",'
                'ISNUMBER(SEARCH("充值巨量千川",AM2)),"巨量千川充值",'
                'ISNUMBER(SEARCH("超",AM2)),"超时赔付",'
                'AND(ISNUMBER(SEARCH("消费者",AM2)),NOT(ISNUMBER(SEARCH("超",AM2))),'
                'NOT(ISNUMBER(SEARCH("运费",AM2)))),"消费者赔付",'
                'F2<>"",VLOOKUP(F2,账单字段对照表!A:B,2,FALSE),F2="",LEFT(AM2,3))')
            if self.formulas != {'AN2': expected_an, 'AO2': '=SUM(Q2:AB2)-R2', 'AP2': expected_ap}:
                raise DataError('template formulas changed or unsupported; refusing to silently change accounting')
            self.mapping = {}
            for scene, category in wb['账单字段对照表'].iter_rows(min_row=2, max_col=2, values_only=True):
                if scene is None:
                    continue
                if scene in self.mapping and self.mapping[scene] != category:
                    raise DataError(f'conflicting template mapping: {scene}')
                self.mapping.setdefault(scene, category)
            self.summary = []
            self.summary_formulas = {}
            self.has_qianchuan = '千川账户' in wb.sheetnames
            self.settlement_columns=[c.value for c in next(wb['结算账单'].iter_rows(max_row=1))]
            if self.settlement_columns[:4]!=['结算时间','订单号','子订单号','结算金额']:
                raise DataError('template settlement column D must be 结算金额')
            summary = wb['核对汇总']
            for row in summary.iter_rows(min_row=2, max_row=43, max_col=4):
                index = row[0].row
                account, category, formula = [c.value for c in row[1:]]
                if index == 2:
                    expected = '=SUMIFS(INDIRECT("\'资金账单 (核对)\'!$AO:$AO"),INDIRECT("\'资金账单 (核对)\'!$AQ:$AQ"),$B2,INDIRECT("\'资金账单 (核对)\'!$AP:$AP"),$C2)'
                elif index in (3,4,5,6):
                    col = {3:'AC',4:'AD',5:'AG',6:'AH'}[index]
                    expected = f'=SUM(INDIRECT("\'资金账单 (核对)\'!${col}:${col}"))'
                else:
                    expected = f'=SUMIFS(INDIRECT("\'资金账单 (核对)\'!$AN:$AN"),INDIRECT("\'资金账单 (核对)\'!$AQ:$AQ"),$B{index},INDIRECT("\'资金账单 (核对)\'!$AP:$AP"),$C{index})'
                if formula != expected:
                    raise DataError(f'unsupported summary formula at D{index}')
                self.summary.append((index, account, category))
                self.summary_formulas[f'D{index}'] = formula
            for coord, expected in [('H2','=SUM(D2:D9)+D31+D37'),('H3','=SUM(INDIRECT("\'结算账单\'!D:D"))'),('H4','=H2-H3')]:
                if summary[coord].value != expected:
                    raise DataError(f'unsupported summary formula at {coord}')
                self.summary_formulas[coord]=expected
            if self.has_qianchuan:
                located={}
                for category,expected in zip(['非赠款消耗','巨量千川充值'],TOTAL_FORMULAS.values()):
                    matches=[row[1].row for row in summary.iter_rows(min_row=2,max_col=4) if row[1].value=='千川账户' and row[2].value==category]
                    if len(matches)!=1:raise DataError(f'千川费用行缺失或重复：{category}')
                    located[f'D{matches[0]}']=expected
                extra=[row[2].value for row in summary.iter_rows(min_row=2,max_col=4) if row[1].value=='千川账户' and row[2].value not in ['非赠款消耗','巨量千川充值']]
                if extra:raise DataError('千川存在未支持的费用行，不能将未知发生额计为零')
                self.qianchuan_cells=located
                for coord, expected in located.items():
                    cell=summary[coord].value
                    formula=getattr(cell,'text',cell)
                    if formula!=expected:raise DataError(f'unsupported summary formula at {coord}')
                    self.summary_formulas[coord]=formula
                    index=int(coord[1:])
                    self.summary.append((index,summary[f'B{index}'].value,summary[f'C{index}'].value))
                if summary['L6'].value!='千川账户' or summary['P6'].value!='=+O6-千川账户!Q4-千川账户!R4':
                    raise DataError('unsupported 千川账户 balance summary')
                self.summary_formulas['P6']=summary['P6'].value
            for index in range(2,7 if self.has_qianchuan else 6):
                for coord,expected in [(f'N{index}',f'=SUMIF(B:D,L{index},D:D)'),(f'O{index}',f'=M{index}+N{index}')]:
                    if summary[coord].value!=expected:raise DataError(f'unsupported summary formula at {coord}')
                    self.summary_formulas[coord]=expected
            self.snapshot = {'template_sha256': self.template_sha256, 'formulas': self.formulas,
                             'summary_formulas': self.summary_formulas, 'mapping': self.mapping.copy(),
                             'array_formulas': [c for c in self.summary_formulas if isinstance(summary[c].value,openpyxl.worksheet.formula.ArrayFormula)],
                             'overrides': {}, 'semantics': 'template-v1-decimal'}
        finally:
            wb.close()
        if overrides:
            config = json.loads(overrides.read_text())
            if config.get('approved') is not True:
                raise DataError('mapping overrides must explicitly declare approved=true')
            self.mapping.update(config['mapping'])
            self.snapshot['overrides'] = config

    def classify(self, values):
        memo = values.get('备注') or ''
        if not isinstance(memo, str):
            memo = str(memo)
        for needle, category in AP_BRANCHES:
            if needle in memo:
                return category, f'memo:{needle}'
        if '消费者' in memo and '超' not in memo and '运费' not in memo:
            return '消费者赔付', 'memo:消费者'
        scene = values.get('动账场景')
        if scene is not None and scene != '':
            return self.mapping.get(scene), f'mapping:{scene}'
        return memo[:3], 'memo:left3'


def process(dataset: Dataset, rules: TemplateRules):
    columns = [Column(k, k, 'decimal' if k in FUND_COLUMNS[16:35] or k in ['动账金额','免佣金额'] else 'datetime' if k in DATE_FIELDS else 'string') for k in FUND_COLUMNS]
    columns += [Column(k,k,t) for k,t in [('结算金额','decimal'),('订单收入','decimal'),('费用类型','string'),('资金账户','string')]]
    processed = Table('processed_funds', '资金账单 (核对)', columns, role='processed')
    dataset.tables.append(processed)
    activity_counts=defaultdict(int)
    for table in dataset.tables.copy():
        if table.role not in ACCOUNTS:
            continue
        time_field = '动账时间' if table.role == 'fund' else '操作时间' if 'deposit' in table.role else '交易时间'
        if table.role == 'steward':
            available = {c.key for c in table.columns}
            time_field = next((c for c in ['交易时间','动账时间','操作时间','创建时间'] if c in available), '交易时间')
        if table.role == 'fund':
            available = {ALIASES.get(c.key, c.key) for c in table.columns}
            missing = set(FUND_COLUMNS) - available
            if missing:
                dataset.issue('missing_fields', 'fund schema does not match template', fields=sorted(missing), table_id=table.table_id)
        for row in target_rows(dataset, table, time_field):
            activity_counts[table.role]+=1
            values = {k: None for k in FUND_COLUMNS}
            account = row.source.get('account') or ACCOUNTS[table.role]
            if table.role == 'fund':
                for key, value in row.values.items():
                    field = ALIASES.get(key, key)
                    if field in values:
                        values[field] = value
            else:
                values.update({'动账时间': row.values.get(time_field),
                               '动帐流水号': next((row.values.get(k) for k in ['操作单号','交易单号','流水号','动账流水号'] if row.values.get(k) is not None), None),
                               '动账金额': next((row.values.get(k) for k in ['操作金额','交易金额','动账金额','金额','动账金额(元)'] if row.values.get(k) is not None), None),
                               '动账方向': row.values.get('动账方向'),
                               '动账场景': row.values.get('操作类型', row.values.get('交易类型', row.values.get('动账场景'))),
                               '订单号': row.values.get('关联订单号'),
                               '备注': row.values.get('备注', row.values.get('交易说明'))})
                if table.role=='steward' and '动账金额(元)' in row.values:
                    # Official wallet CSV amounts are signed already. Convert
                    # to the template's absolute amount + direction convention,
                    # so its retained AN formula never reverses a debit twice.
                    try:
                        wallet_amount=decimal(values['动账金额'],blank_zero=False)
                        direction=str(row.values.get('账户方向') or '').strip()
                        if direction not in ['收入','支出'] or (direction=='支出' and wallet_amount>0) or (direction=='收入' and wallet_amount<0):
                            raise DataError('wallet amount and account direction disagree')
                        values['动账金额']=abs(wallet_amount)
                        values['动账方向']='出账' if direction=='支出' else '入账'
                        values['子订单号']=row.values.get('关联子订单号')
                    except DataError as exc:
                        dataset.issue('invalid_amount',str(exc),source=row.source)
                        continue
                # Normalize padding in the derived scene only; raw data remains untouched.
                if isinstance(values['动账场景'], str):
                    values['动账场景'] = values['动账场景'].strip()
                # Verified deposit CSVs prefix cells with tabs and format CNY
                # amounts as +¥100.00 / -¥100.00. Normalize only this adapter;
                # original source values remain intact in the raw table.
                if 'deposit' in table.role and isinstance(values['动账金额'],str):
                    money=values['动账金额'].strip()
                    if re.fullmatch(r'[+-]?[¥￥](?:\d{1,3}(?:,\d{3})+|\d+)(?:\.\d{1,2})?',money):
                        values['动账金额']=money.replace('¥','').replace('￥','').replace(',','')
                if '待确认' in account:
                    dataset.issue('account_unknown', 'deposit account type needs explicit source metadata', source=row.source)
            try:
                amount = decimal(values['动账金额'], blank_zero=False)
                signed = -amount if values['动账方向'] == '出账' else amount
                income = sum((decimal(values[k]) for k in FUND_COLUMNS[16:28]), Decimal(0)) - decimal(values['运费实付'])
            except DataError as exc:
                dataset.issue('invalid_amount', str(exc), source=row.source)
                continue
            category, matched = rules.classify(values)
            if category is None or category == '':
                dataset.issue('unmapped_scene', 'expense category could not be determined', source=row.source,
                              scene=values['动账场景'], amount=str(signed))
            values.update({'结算金额': signed, '订单收入': income, '费用类型': category, '资金账户': account})
            processed.rows.append(Row(values, {**row.source, 'source_row_id': row.row_id, 'classification_rule': matched}, row.row_id))
    sums = defaultdict(Decimal)
    account_raw = defaultdict(Decimal)
    expense_totals = defaultdict(Decimal)
    for row in processed.rows:
        v = row.values
        key = (v['资金账户'], v['费用类型'])
        sums[(*key,'AN')] += v['结算金额']
        sums[(*key,'AO')] += v['订单收入']
        account_raw[v['资金账户']] += v['结算金额']
        for field in ['平台服务费','佣金','招商服务费','站外推广费']:
            expense_totals[field] += decimal(v[field])
        for field in ['服务商佣金','渠道分成','其他分成']:
            if decimal(v[field]) != 0:
                dataset.issue('template_fee_not_listed', f'nonzero {field} has no dedicated template summary row', source=row.source)
    qianchuan=prepare_qianchuan(dataset,rules.has_qianchuan)
    summary_records = []
    amounts = {}
    keys = {(a,c) for _,a,c in rules.summary}
    for index, account, category in rules.summary:
        amount = expense_totals[category] if index in (3,4,5,6) else sums[(account, category, 'AO' if index == 2 else 'AN')]
        if account=='千川账户':amount=qianchuan['consumption' if category=='非赠款消耗' else 'recharge'] if qianchuan else None
        amounts[index] = amount
        summary_records.append({'店铺编码': dataset.context.get('platform_shop_id'), '资金账户': account,
                                '费用类型': category, '金额': amount, '模板行号': index})
    for row in processed.rows:
        pair = (row.values['资金账户'], row.values['费用类型'])
        if pair[1] and pair not in keys:
            dataset.issue('unrouted_category', 'classified row has no template summary item', source=row.source,
                          account=pair[0], category=pair[1])
    dataset.tables.append(Table.from_records('fee_summary','费用分类汇总',summary_records,types={'金额':'decimal','模板行号':'integer'}))
    settled = Decimal(0)
    # Keep the source's display order and spelling (e.g. 店铺券 / APP渠道),
    # while fixing the first four columns for the template's H3 dependency.
    settlement_fields=list(dict.fromkeys([*rules.settlement_columns[:4],*(c.key for t in dataset.table('settlement') for c in t.columns)]))
    canonical={ALIASES.get(key,key) for key in settlement_fields}
    for key in rules.settlement_columns[4:]:
        if ALIASES.get(key,key) not in canonical:
            settlement_fields.append(key);canonical.add(ALIASES.get(key,key))
    monthly=Table('settlement_monthly','结算账单',
                  [Column(k,k,'decimal' if k in AMOUNT_FIELDS else 'datetime' if k in DATE_FIELDS else 'string') for k in settlement_fields],role='settlement_monthly')
    for table in dataset.table('settlement'):
        for row in target_rows(dataset, table, '结算时间'):
            try:
                amount=decimal(row.values.get('结算金额'), blank_zero=False)
                settled += amount
                monthly.rows.append(Row({**row.values,'结算金额':amount},{**row.source,'source_row_id':row.row_id},row.row_id))
            except DataError as exc:
                dataset.issue('invalid_settlement', str(exc), source=row.source)
    # The template's H3 sums column D of this month-filtered, merged view.
    # Keep every original source sheet separately, with its original columns.
    dataset.tables.append(monthly)
    goods = sum((amounts[r] for r in range(2,10)), Decimal(0)) + amounts[31] + amounts[37]
    dataset.checks.append({'项目':'结算金额核对','计算值':goods,'实际值':settled,'差异':goods-settled,
                           '状态':'一致' if goods == settled else '存在差异'})
    accounts = []
    for account in ['货款账户','管家账户','基础保证金账户','体验保证金账户']:
        roles={role for role,name in ACCOUNTS.items() if name==account}
        receipts=[c for c in dataset.context.get('collection',[]) if c.get('role') in roles]
        source_tables=[t for t in dataset.tables if t.role in roles]
        normal_empty=(bool(receipts) and all(c.get('status') in ['no_activity','not_opened'] for c in receipts)
                      or bool(source_tables))
        normal_empty=normal_empty and not any(activity_counts[role] for role in roles) and not any(c.get('status')=='error' for c in receipts)
        known=(any(t.role in roles for t in dataset.tables) or normal_empty) and not any(c.get('status')=='error' for c in receipts)
        movement = sum((r['金额'] for r in summary_records if r['资金账户']==account), Decimal(0))
        gap = movement - account_raw[account]
        dataset.checks.append({'项目':f'{account}流水覆盖','计算值':movement if known else None,'实际值':account_raw[account] if known else None,'差异':gap if known else None,
                               '状态':('一致' if gap == 0 else '存在差异') if known else '待确认'})
        if not known:dataset.issue('account_source_missing',f'{account}: ledger unavailable or incompletely collected',account=account)
        evidence = [b for b in dataset.balances if b.get('account') == account]
        valid = [b for b in evidence if b.get('period') == dataset.context['month'] and b.get('kind') == 'monthly'
                 and b.get('opening') is not None and b.get('closing') is not None and b.get('source')]
        if len(valid) != 1:
            if not normal_empty:dataset.issue('balance_missing', f'{account}: need one verified monthly opening/closing balance', account=account)
            accounts.append({'资金账户':account,'期初余额':None,'本期发生额':movement if known else None,'计算期末余额':None,'实际期末余额':None,'差异':None,'状态':'无余额记录' if normal_empty else '待确认'})
            continue
        b = valid[0]
        opening, actual = decimal(b['opening']), decimal(b['closing'])
        expected = opening + movement
        accounts.append({'资金账户':account,'期初余额':opening,'本期发生额':movement,'计算期末余额':expected,
                         '实际期末余额':actual,'差异':expected-actual,'状态':'一致' if expected==actual else '存在差异'})
        dataset.checks.append({'项目':f'{account}余额','计算值':expected,'实际值':actual,'差异':expected-actual,
                               '状态':'一致' if expected == actual else '存在差异'})
    if rules.has_qianchuan:
        movement=qianchuan['consumption']+qianchuan['recharge'] if qianchuan else None
        actual=qianchuan['grant']+qianchuan['cash'] if qianchuan and qianchuan['grant'] is not None and qianchuan['cash'] is not None else None
        valid=[b for b in dataset.balances if b.get('account')=='千川账户' and b.get('period')==dataset.context['month']
               and b.get('kind')=='monthly' and b.get('opening') is not None and b.get('source')]
        opening=decimal(valid[0]['opening']) if len(valid)==1 else None
        expected=opening+movement if opening is not None and movement is not None else None
        gap=expected-actual if expected is not None and actual is not None else None
        normal_missing_balance=qianchuan is not None and actual is None
        denied=dataset.context.get('qianchuan_access')=='permission_denied'
        if opening is None and not normal_missing_balance and not denied:dataset.issue('balance_missing','千川账户: need one verified monthly opening balance',account='千川账户')
        state='无月末余额记录' if normal_missing_balance else '待确认' if gap is None else '一致' if gap==0 else '存在差异'
        if denied:state='无权限（数据留空）'
        accounts.append({'资金账户':'千川账户','期初余额':opening,'本期发生额':movement,'计算期末余额':expected,
                         '实际期末余额':actual,'差异':gap,'状态':state})
        dataset.checks.append({'项目':'千川账户余额','计算值':expected,'实际值':actual,'差异':gap,
                               '状态':state})
    dataset.tables.append(Table.from_records('account_balances','账户余额核对',accounts,
                          types={k:'decimal' for k in ['期初余额','本期发生额','计算期末余额','实际期末余额','差异']}))
    for role in ('fund','settlement'):
        if not dataset.table(role) and not any(c.get('role')==role and c.get('status')=='no_activity' for c in dataset.context.get('collection',[])):
            dataset.issue('source_missing', f'missing required source: {role}')
    dataset.context['balance_validation']=[c['项目']+'：'+c['状态'] for c in dataset.checks if '余额' in c['项目'] and c['状态']!='一致']
    dataset.context['fee_source_errors']=[i['message'] for i in dataset.issues if i['code'] in ['source_missing','invalid_amount','invalid_settlement','missing_fields','row_count_mismatch','collection_failed']]
    dataset.context['rule_snapshot'] = rules.snapshot
    dataset.context['conclusion'] = ('待确认' if any(i['severity']=='blocking' for i in dataset.issues)
                                     else '存在差异' if any(c['差异'] is not None and c['差异'] != 0 for c in dataset.checks)
                                     else '无数据' if dataset.context.get('collection') and all(c.get('status') in ['no_activity','not_opened'] for c in dataset.context['collection'])
                                     else '核对完成（部分余额无记录）' if any(c['差异'] is None for c in dataset.checks) else '完全一致')
    return processed
