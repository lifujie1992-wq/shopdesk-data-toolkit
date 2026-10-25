from __future__ import annotations

import csv
from datetime import datetime
import io
from pathlib import Path, PurePosixPath
import posixpath
import shutil
import zipfile
import xml.etree.ElementTree as ET

import openpyxl

from .model import Column, DataError, Dataset, Row, Table, month_bounds, parse_time, safe_name, sha256
from .schema import AMOUNT_FIELDS, DATE_FIELDS, ID_FIELDS, detect_role

SKIP_SHEETS = {'资金账单 (核对)', '核对汇总', '账单字段对照表', '核对结论', '差异明细'}


def _column(name):
    kind = 'decimal' if name in AMOUNT_FIELDS else 'datetime' if name in DATE_FIELDS else 'string'
    return Column(name, name, kind)


def _read_table(dataset, name, rows, artifact, role=None, account=None, entity_id=None, entity_name=None):
    # Some official CSVs have preambles. Recognize a header in the first 30 rows.
    header = None
    table = None
    for index, values in enumerate(rows, 1):
        if header is None:
            if index > 30:
                break
            fields = [str(v).strip() if v is not None else '' for v in values]
            found = detect_role(fields, name)
            if not found and role in ['tax','steward'] and len([f for f in fields if f])>=2:
                found=role
            if not found:
                continue
            header = fields
            if found == 'qianchuan':
                # A completed monthly workbook has sparse Q/R balance annotations.
                # They are not extra columns of the downloaded daily report.
                header = header[:14]
            while header and not header[-1]:
                header.pop()
            if any(not x for x in header) or len(set(header)) != len(header):
                raise DataError(f'{name}: duplicate or blank column names')
            if role in ['fund','settlement','base_deposit','experience_deposit','deposit','qianchuan'] and role != found and not ('deposit' in role and 'deposit' in found):
                raise DataError(f'{name}: source schema {found} conflicts with declared role {role}')
            role = role or found
            table = Table(f'{role}_{artifact["artifact_id"][:8]}_{len(dataset.tables)}', name,
                          [_column(h) for h in header], role=role)
            dataset.tables.append(table)
            continue
        if not any(v is not None and v != '' for v in values):
            continue
        if role=='qianchuan':
            values=values[:len(header)]
            if not any(v is not None and v != '' for v in values):continue
        if len(values) > len(header) and any(v is not None for v in values[len(header):]):
            dataset.issue('extra_cells', 'data extends beyond the header', source={'sheet': name, 'row': index})
        record = {}
        for col, value in zip(table.columns, list(values) + [None] * len(header)):
            if col.name in ID_FIELDS and value is not None:
                if isinstance(value, (int, float)) and len(str(int(value))) > 15:
                    dataset.issue('numeric_identifier', 'long identifier stored as a number; source precision may be lost',
                                  source={'artifact_id': artifact['artifact_id'], 'sheet': name, 'row': index},
                                  field=col.name)
                value = str(value)
            record[col.key] = value
        source = {'artifact_id': artifact['artifact_id'], 'sheet': name, 'row': index,
                  'account': account, 'entity_id': entity_id, 'entity_name': entity_name, 'role': role}
        if role=='settlement' and str(record.get('结算金额') or '').strip()=='收入+支出' and all(not str(record.get(k) or '').strip() for k in ['结算时间','订单号','子订单号']):
            source['kind']='column_note'
        table.rows.append(Row(record, source))
    return table


def _media_from_xlsx(path, out_dir, artifact_id):
    ns = {'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
          'p':'http://schemas.openxmlformats.org/package/2006/relationships',
          'xdr':'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
          'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
    result = []
    def resolve(base, target):
        return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(base), target))
    def rels(z, base):
        rpath = posixpath.join(posixpath.dirname(base), '_rels', posixpath.basename(base) + '.rels')
        if rpath not in z.namelist():
            return {}
        return {n.attrib['Id']: resolve(base, n.attrib['Target'])
                for n in ET.fromstring(z.read(rpath)) if n.attrib.get('TargetMode') != 'External'}
    with zipfile.ZipFile(path) as z:
        wb = ET.fromstring(z.read('xl/workbook.xml'))
        wrels = rels(z, 'xl/workbook.xml')
        for sheet in wb.findall('m:sheets/m:sheet', ns):
            sheet_path = wrels[sheet.attrib['{' + ns['r'] + '}id']]
            srels = rels(z, sheet_path)
            root = ET.fromstring(z.read(sheet_path))
            for drawing in root.findall('m:drawing', ns):
                drawing_path = srels[drawing.attrib['{' + ns['r'] + '}id']]
                drels = rels(z, drawing_path)
                for anchor in ET.fromstring(z.read(drawing_path)):
                    for blip in anchor.findall('.//a:blip', ns):
                        image_path = drels.get(blip.attrib.get('{' + ns['r'] + '}embed'))
                        if not image_path:
                            continue
                        raw = z.read(image_path)
                        output = out_dir / f'{artifact_id[:12]}_{safe_name(sheet.attrib["name"])}_{len(result)+1}{PurePosixPath(image_path).suffix}'
                        output.parent.mkdir(parents=True, exist_ok=True)
                        output.write_bytes(raw)
                        result.append({'path': str(output), 'sheet': sheet.attrib['name'],
                                       'artifact_id': artifact_id, 'sha256': sha256(output),
                                       'source_media': image_path, 'balance_as_of': None,
                                       'anchor_xml': ET.tostring(anchor, encoding='unicode'),
                                       'anchor_from': {n.tag.rsplit('}',1)[-1]:int(n.text) for n in anchor.find('xdr:from',ns)} if anchor.find('xdr:from',ns) is not None else None,
                                       'note': 'original screenshot; no OCR balance inferred'})
    return result


def import_file(dataset: Dataset, path: Path, run_dir: Path, *, role=None, account=None, entity_id=None, entity_name=None):
    path = path.resolve()
    if not path.is_file():
        raise DataError(f'input not found: {path}')
    digest = sha256(path)
    if any(a['sha256'] == digest and a.get('role') == role and a.get('account') == account and a.get('entity_id') == entity_id and a.get('entity_name') == entity_name
           for a in dataset.artifacts):
        return
    raw_dir = run_dir / 'raw'
    raw_dir.mkdir(parents=True, exist_ok=True)
    target = raw_dir / f'{digest[:12]}_{safe_name(path.name)}'
    if path != target.resolve():
        shutil.copy2(path, target)
    artifact = {'artifact_id': digest, 'path': str(target), 'source_name': path.name,
                'sha256': digest, 'size': target.stat().st_size, 'role': role,
                'account': account, 'entity_id': entity_id, 'entity_name':entity_name}
    dataset.artifacts.append(artifact)
    if path.suffix.lower() in {'.xlsx', '.xlsm'}:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            generated_monthly=('核对结论' in wb.sheetnames and '资金账单 (核对)' in wb.sheetnames
                               and any(n.startswith('结算账单原始') for n in wb.sheetnames))
            for sheet in wb:
                if sheet.title in SKIP_SHEETS:
                    continue
                if generated_monthly and sheet.title=='结算账单':
                    # Our formula view merges the raw settlement sheets. When
                    # reimporting a generated workbook, count the raw sources once.
                    continue
                parsed = _read_table(dataset, sheet.title, sheet.iter_rows(values_only=True), artifact,
                                     role, account, entity_id, entity_name)
                if parsed is None:
                    # Balance/image sheets can contain sparse handwritten reference cells.
                    # Preserve them for online-document sync without treating them as transactions.
                    refs=[]
                    for cells in sheet.iter_rows():
                        for cell in cells:
                            if cell.value is not None:
                                refs.append(Row({'坐标':cell.coordinate,'原始值':cell.value,'字段类型':cell.data_type},
                                                {'artifact_id':artifact['artifact_id'],'sheet':sheet.title,'row':cell.row,'column':cell.column}))
                    if refs:
                        dataset.tables.append(Table(f'reference_{digest[:8]}_{len(dataset.tables)}',
                            sheet.title+'原始记录',[Column('坐标','坐标'),Column('原始值','原始值'),Column('字段类型','字段类型')],
                            rows=refs,role='reference'))
        finally:
            wb.close()
        images = _media_from_xlsx(path, run_dir / 'screenshots', digest)
        if images:
            dataset.context.setdefault('images', []).extend(images)
    elif path.suffix.lower() in {'.csv', '.tsv'}:
        raw = path.read_bytes()
        for encoding in ('utf-8-sig', 'gb18030'):
            try:
                text = raw.decode(encoding)
                break
            except UnicodeDecodeError:
                pass
        else:
            raise DataError(f'unsupported CSV encoding: {path.name}')
        delimiter = '\t' if path.suffix.lower() == '.tsv' else ','
        if '\t' in text.splitlines()[0] and ',' not in text.splitlines()[0]:
            delimiter = '\t'
        _read_table(dataset, path.stem, csv.reader(io.StringIO(text), delimiter=delimiter),
                    artifact, role, account, entity_id, entity_name)
    elif path.suffix.lower() == '.zip':
        expanded = run_dir / 'unpacked' / digest[:12]
        with zipfile.ZipFile(path) as z:
            for info in z.infolist():
                if info.is_dir():
                    continue
                relative = PurePosixPath(info.filename)
                if relative.is_absolute() or '..' in relative.parts:
                    raise DataError('unsafe ZIP member path')
                if info.file_size > 512 * 1024 * 1024:
                    raise DataError('ZIP member exceeds 512MB')
                if relative.suffix.lower() not in {'.xlsx','.csv','.tsv'}:
                    continue
                out = expanded.joinpath(*relative.parts)
                out.parent.mkdir(parents=True, exist_ok=True)
                with z.open(info) as src, out.open('wb') as dst:
                    shutil.copyfileobj(src, dst)
                import_file(dataset, out, run_dir, role=role, account=account, entity_id=entity_id, entity_name=entity_name)
    else:
        raise DataError(f'unsupported input format: {path.suffix}')


def target_rows(dataset, table, time_field):
    start, end = month_bounds(dataset.context['month'])
    for row in table.rows:
        if row.source.get('kind')=='column_note':continue
        try:
            at = parse_time(row.values.get(time_field))
        except DataError as exc:
            dataset.issue('invalid_date', str(exc), source=row.source)
            continue
        if at.tzinfo is not None:
            dataset.issue('timezone', 'source timestamp needs an explicit timezone conversion', source=row.source)
            continue
        if start <= at < end:
            yield row
        else:
            dataset.issue('outside_month', 'source row is outside target month; retained in raw table, excluded from calculations',
                          severity='info', source=row.source)
