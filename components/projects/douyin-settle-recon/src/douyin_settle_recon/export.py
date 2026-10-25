from __future__ import annotations

import csv
from datetime import date, datetime, timezone, timedelta
from decimal import Decimal
import gzip
import json
import re
from pathlib import Path
import shutil
import sqlite3

import xlsxwriter
from openpyxl.formula.translate import Translator

from .model import Dataset, Table, DataError, decimal, dump_json, json_value, safe_name, sha256
from .workbook_layout import ReportLayout, sheet_name


def _scalar(value):
    return json_value(value) if isinstance(value, (Decimal,date,datetime)) else value


def _cell(ws, row, col, value, column, formats):
    if value is None:
        return
    if isinstance(value, (datetime,date)):
        ws.write_datetime(row,col,value,formats['date'])
    elif isinstance(value,bool):
        ws.write_boolean(row,col,value)
    elif isinstance(value,(int,float,Decimal)):
        ws.write_number(row,col,float(value),formats['amount'] if column.type=='decimal' else None)
    else:
        # Explicit string writing prevents formula/URL interpretation and preserves long IDs.
        text=str(value)
        if len(text)>32767:raise DataError('a cell exceeds Excel 32767-character capacity; raw file remains available')
        ws.write_string(row,col,text,formats['text'] if column.type=='string' else None)


def export(dataset: Dataset, run_dir: Path):
    run_dir.mkdir(parents=True,exist_ok=True)
    context = dataset.context
    snapshot=context['rule_snapshot']
    mapping={**snapshot['mapping'],**snapshot.get('overrides',{}).get('mapping',{})}
    dataset.tables.append(Table.from_records('bill_mapping','账单字段对照表',
        [{'动账场景':scene,'费用类型':category} for scene,category in mapping.items()],role='mapping',fields=['动账场景','费用类型']))
    layout = ReportLayout(run_dir/'rules/template.xlsx', dataset.artifacts)
    # Tables are the canonical output shared by Excel, SQLite, JSONL, CSV and future online adapters.
    dataset.tables.extend([
        Table.from_records('checks','核对检查',dataset.checks,types={k:'decimal' for k in ['计算值','实际值','差异']}),
        Table.from_records('issues','异常明细',[{**i,'source':json.dumps(i.get('source',{}),ensure_ascii=False)} for i in dataset.issues],fields=['code','severity','message','source','role','account','category','scene','amount','fields','table_id','expected','actual']),
        Table.from_records('balance_evidence','余额依据',[{**b,'source':json.dumps(b.get('source'),ensure_ascii=False)} for b in dataset.balances],fields=['account','period','kind','opening','closing','movement','expected_count','source'],types={'opening':'decimal','closing':'decimal','movement':'decimal','expected_count':'integer'}),
    ])
    db_path = run_dir/'tables.sqlite'
    db = sqlite3.connect(db_path)
    db.executescript('''CREATE TABLE context(key TEXT PRIMARY KEY,value TEXT NOT NULL);
        CREATE TABLE tables(table_id TEXT PRIMARY KEY,schema_json TEXT NOT NULL);
        CREATE TABLE rows(table_id TEXT NOT NULL,ordinal INTEGER NOT NULL,row_id TEXT NOT NULL,values_json TEXT NOT NULL,source_json TEXT NOT NULL,
                          PRIMARY KEY(table_id,ordinal));
        CREATE INDEX row_identity ON rows(table_id,row_id);''')
    db.execute('INSERT INTO context VALUES (?,?)',('run',json.dumps({k:v for k,v in context.items() if k not in ['images','rule_snapshot']},ensure_ascii=False,default=json_value)))
    schemas = []
    prefix = 'Y-' if context.get('platform', 'douyin') == 'douyin' else ''
    output_name = f'{context["month"].replace("-", "")}【抖音】{prefix}{safe_name(context["shop_name"])}.xlsx'
    path = run_dir/output_name
    temp = run_dir/('.'+output_name+'.tmp')
    for table in dataset.tables:
        schema=table.schema();schema['excel_sheet']=_table_sheet(table)
        schema['context']={k:context.get(k) for k in ['run_id','month','shop_id','shopdesk_shop_uuid','platform_shop_id','shop_name']}
        schemas.append(schema)
        directory=run_dir/'tables'/safe_name(table.table_id);directory.mkdir(parents=True,exist_ok=True)
        dump_json(directory/'schema.json',schema)
        db.execute('INSERT INTO tables VALUES (?,?)',(table.table_id,json.dumps(schema,ensure_ascii=False,default=json_value)))
        with (directory/'rows.jsonl').open('w',encoding='utf-8') as js,(directory/'rows.csv').open('w',encoding='utf-8-sig',newline='') as cf:
            writer=csv.writer(cf);writer.writerow([c.name for c in table.columns])
            for index,row in enumerate(table.rows,1):
                values={c.key:(str(row.values[c.key]) if c.type=='decimal' and row.values.get(c.key) is not None else _scalar(row.values.get(c.key))) for c in table.columns}
                js.write(json.dumps({'row_id':row.row_id,'values':values,'source':row.source},ensure_ascii=False,default=json_value)+'\n')
                writer.writerow([values[c.key] for c in table.columns])
                db.execute('INSERT INTO rows VALUES (?,?,?,?,?)',(table.table_id,index,row.row_id,json.dumps(values,ensure_ascii=False),json.dumps(row.source,ensure_ascii=False,default=json_value)))
    image_rows = _write_report(dataset, temp, layout)
    layout.finish(temp, image_rows)
    temp.replace(path)
    db.commit();db.close()
    from .report_validation import validate_report
    validation=validate_report(path,context,snapshot)
    context['upload_ready']=validation['ready']
    context['partial_accounts']=validation['partial_accounts']
    dump_json(run_dir/'report-validation.json',validation)
    dump_json(run_dir/'qianchuan-result.json',context.get('qianchuan_result',{}))
    final_files={path:{'sha256':sha256(path),'size':path.stat().st_size}}
    gzip_path=path.with_suffix('.xlsx.gz')
    if validation['ready']:
        with path.open('rb') as src,gzip_path.open('wb') as dst:
            with gzip.GzipFile(filename='',mode='wb',fileobj=dst,mtime=0) as zipped:shutil.copyfileobj(src,zipped)
        final_files[gzip_path]={'sha256':sha256(gzip_path),'size':gzip_path.stat().st_size}
    dump_json(run_dir/'tables.json',{'schema_version':'1.0','context':{k:v for k,v in context.items() if k not in ['images','rule_snapshot']},'tables':schemas,
                                    'online_document_contract':{'column_order':'schema.columns order','decimal_encoding':'base-10 string','identifier_encoding':'string','row_identity':'row_id','images':'separate artifacts; not OCR values'}})
    dump_json(run_dir/'rules.json',context.get('rule_snapshot',{}))
    dump_json(run_dir/'result.json',{'context':{k:v for k,v in context.items() if k not in ['images','rule_snapshot']},'checks':dataset.checks,'issues':dataset.issues,'table_counts':{t.table_id:len(t.rows) for t in dataset.tables}})
    # All paths in the handoff manifest are relative, so copying a run keeps it portable.
    artifacts=[]
    for file in sorted(run_dir.rglob('*')):
        if file.is_file() and file.name not in {'manifest.json','.run.lock'}:
            info=final_files.get(file) or {'sha256':sha256(file),'size':file.stat().st_size}
            artifacts.append({'artifact_id':info['sha256'],'path':str(file.relative_to(run_dir)),**info})
    dump_json(run_dir/'manifest.json',{'schema_version':'1.0','shop':{k:context.get(k) for k in ['platform','shop_id','shopdesk_shop_uuid','platform_shop_id','shop_name','shop_identity','agent_id']},
                                      'month':context['month'],'run_id':context['run_id'],'conclusion':context['conclusion'],
                                      **({'reconciliation_date':context['reconciliation_date']} if context.get('reconciliation_date') else {}),
                                      'artifacts':artifacts,'source_artifacts':[{**a,'path':str(Path(a['path']).relative_to(run_dir))} for a in dataset.artifacts],'screenshots':[{**a,'path':str(Path(a['path']).relative_to(run_dir))} for a in image_rows],
                                      'report':str(path.relative_to(run_dir)),'upload_ready':validation['ready'],
                                      'partial_accounts':validation['partial_accounts'],
                                      'report_gzip':str(gzip_path.relative_to(run_dir)) if validation['ready'] else None,
                                      'tables':'tables.json','database':'tables.sqlite'})
    return path


def _table_sheet(table):
    if table.table_id=='processed_funds':return '资金账单 (核对)'
    if table.table_id=='settlement_monthly':return '结算账单'
    if table.table_id=='bill_mapping':return '账单字段对照表'
    if table.role=='qianchuan_monthly':return '千川账户'
    if table.role=='fund':return '资金账单'
    # Wallet ledgers remain canonical inputs for reconciliation and structured
    # exports. The visible wallet sheet contains screenshots and balances only.
    if table.role=='steward':return None
    if table.role in ['deposit','base_deposit','experience_deposit']:return '保证金账户'
    if table.role=='reference' and table.rows:
        name=sheet_name(table.rows[0].source.get('sheet',''))
        return name if name in ['资金账单','结算账单','管家账户','保证金账户','余额截图'] else None
    # These are retained in SQLite/CSV/JSONL, without adding report worksheets.
    return None


def _write_report(dataset, path, layout):
    wb=xlsxwriter.Workbook(path,{'constant_memory':True,'strings_to_formulas':False,'strings_to_urls':False})
    wb.set_calc_mode('auto')
    sheets={name:wb.add_worksheet(name) for name in layout.sheetnames}
    sheets['核对汇总'].activate()
    formats={'amount':wb.add_format({'num_format':'0.00########'}),
             'date':wb.add_format({'num_format':'yyyy-mm-dd hh:mm:ss'}),
             'text':wb.add_format({'num_format':'@'})}
    wallet_col=8  # Keep the total at J7 for existing balance-check references.
    wallet_grid=_wallet_balance_grid(dataset,wallet_col)
    wallet_header=wb.add_format({'bold':True,'bg_color':'#DCE6F1','border':1})
    wallet_total=wb.add_format({'bold':True,'font_size':14,'num_format':'0.00','bg_color':'#E2EFDA'})
    wallet_amount=wb.add_format({'num_format':'0.00','border':1})
    wallet_text=wb.add_format({'border':1})
    def write_wallet_row(sheet,rn):
        for col,value in sorted(wallet_grid.get(rn,{}).items()):
            fmt=wallet_header if rn in [0,9] else wallet_total if rn==6 else wallet_amount if col==wallet_col+1 else wallet_text
            if isinstance(value,tuple):sheet.write_formula(rn,col,value[0],fmt,float(value[1]) if isinstance(value[1],Decimal) else value[1])
            elif isinstance(value,Decimal):sheet.write_number(rn,col,float(value),fmt)
            elif value is not None:sheet.write_string(rn,col,str(value),fmt)
    formula_columns={key:Translator(dataset.context['rule_snapshot']['formulas'][coord],origin=coord)
                     for key,coord in [('结算金额','AN2'),('订单收入','AO2'),('费用类型','AP2')]}
    for name,sheet in sheets.items():
        if name=='核对汇总':
            _write_template_summary(sheet,dataset,formats,layout)
            continue
        if name=='千川账户':
            _write_qianchuan(sheet,dataset,formats)
            continue
        if name=='管家账户':
            sheet.set_column(wallet_col,wallet_col,36)
            sheet.set_column(wallet_col+1,wallet_col+1,16)
            sheet.set_column(wallet_col+2,wallet_col+2,14)
            sheet.set_column(wallet_col+3,wallet_col+3,25)
            if wallet_grid:
                for rn in sorted(wallet_grid):write_wallet_row(sheet,rn)
            else:
                # Importing a legacy workbook may provide only a manually
                # verified total. Never copy its raw ledger cells back here.
                legacy=layout.cells(name);total=legacy.get('J7')
                sheet.write_string(0,wallet_col,'管家账户余额汇总',wallet_header)
                sheet.write_string(6,wallet_col,'总余额（元）',wallet_total)
                if isinstance(total,tuple):sheet.write_formula(6,wallet_col+1,total[0],wallet_total,total[1] if total[1] is not None else '')
                elif total is not None and legacy.get('I7')=='总余额（元）':sheet.write_number(6,wallet_col+1,float(decimal(total)),wallet_total)
                else:sheet.write_formula(6,wallet_col+1,'=NA()',wallet_total,'#N/A')
                labels=['账户主体','总余额（元）','账户状态','截图时间（北京时间）']
                sheet.write_row(9,wallet_col,labels,wallet_header)
                if [legacy.get(f'{col}10') for col in ['I','J','K','L']]==labels:
                    for coord,value in sorted(legacy.items(),key=lambda item:xlsxwriter.utility.xl_cell_to_rowcol(item[0])):
                        rn,col=xlsxwriter.utility.xl_cell_to_rowcol(coord)
                        if rn<10 or not wallet_col<=col<=wallet_col+3:continue
                        if isinstance(value,tuple):sheet.write_formula(rn,col,value[0],wallet_amount if col==wallet_col+1 else wallet_text,value[1] if value[1] is not None else '')
                        else:_cell(sheet,rn,col,value,ColumnStub('decimal' if col==wallet_col+1 else 'string'),{'amount':wallet_amount,'text':wallet_text,'date':formats['date']})
            continue
        tables=[t for t in dataset.tables if _table_sheet(t)==name and t.role!='reference']
        if tables:
            columns=list({c.key:c for t in tables for c in t.columns}.values())
            for col,c in enumerate(columns):sheet.write_string(0,col,c.name)
            index=0
            for table in tables:
                for row in table.rows:
                    index+=1
                    if index>1_048_575:raise DataError('Excel row capacity exceeded')
                    for col,column in enumerate(columns):
                        # The template appends account movements using amount,
                        # scene and memo. Full account identifiers/timestamps
                        # remain in canonical raw table rows.
                        if (table.table_id=='processed_funds'
                            and row.source.get('role') in ['deposit','base_deposit','experience_deposit','steward']
                            and column.key in ['动账时间','动帐流水号','订单号']):continue
                        value=row.values.get(column.key)
                        if table.table_id=='processed_funds' and column.key in formula_columns:
                            coord=xlsxwriter.utility.xl_rowcol_to_cell(index,col)
                            formula=formula_columns[column.key].translate_formula(coord)
                            cached='#N/A' if value is None and column.key=='费用类型' else float(value) if isinstance(value,Decimal) else value
                            sheet.write_formula(index,col,formula,formats['amount'] if column.type=='decimal' else None,cached)
                        else:
                            if table.table_id in ['processed_funds','settlement_monthly'] and column.type=='decimal' and isinstance(value,str) and value.strip():
                                try:value=decimal(value,blank_zero=False)
                                except DataError:pass
                            _cell(sheet,index,col,value,column,formats)
        else:
            # Sparse annotations are kept at their original coordinates, with
            # their source formulas rather than the data_only cached value.
            reference_cells=dict(layout.cells(name))
            verified=_fund_month_balance(dataset)
            if name=='余额截图' and verified is not None:reference_cells['F27']=decimal(verified['closing'])
            for coord,value in sorted(reference_cells.items(),key=lambda item:xlsxwriter.utility.xl_cell_to_rowcol(item[0])):
                rn,col=xlsxwriter.utility.xl_cell_to_rowcol(coord)
                if isinstance(value,tuple):sheet.write_formula(rn,col,value[0],None,value[1] if value[1] is not None else '')
                else:_cell(sheet,rn,col,value,ColumnStub('string' if isinstance(value,str) else 'decimal'),formats)
    images=[];counts={};next_rows={}
    for number,original in enumerate(dataset.context.get('images',[]),1):
        item=dict(original);name=_image_sheet(item)
        if name=='千川账户' and dataset.context.get('qianchuan_access')=='permission_denied':
            images.append({**item,'excel_sheet':None,'omitted_reason':'qianchuan_permission_denied'})
            continue
        if name=='千川账户' and item.get('query',{}).get('total')==0:
            images.append({**item,'excel_sheet':None,'omitted_reason':'empty_qianchuan_statement'})
            continue
        if name is None:
            images.append({**item,'excel_sheet':None})
            continue
        if name not in sheets:raise DataError(f'screenshot has no template destination: {item.get("sheet")}')
        count=counts.get(name,0);counts[name]=count+1
        anchor=item.get('anchor_from')
        if anchor:
            rn,col=anchor['row'],anchor['col']
            options={'x_offset':anchor.get('colOff',0)/9525,'y_offset':anchor.get('rowOff',0)/9525}
        else:
            # New account captures follow the example: deposits on the right,
            # steward/balance evidence stacked on the corresponding account page.
            rn,col=(count*19,8) if name=='保证金账户' else (count*16,0)
            options={}
        from PIL import Image
        with Image.open(item['path']) as image:
            width=510 if name=='管家账户' and wallet_grid and not anchor else 600
            scale=min(1,width/image.width);height=image.height*scale
        if not anchor:
            rn=next_rows.get(name,0)
            if name=='千川账户':rn=max(rn,7);col=max(col,15)
            next_rows[name]=rn+int(height/20)+2
            tables=[t for t in dataset.tables if _table_sheet(t)==name and t.role!='reference']
            if tables:col=max(col,max(len(t.columns) for t in tables)+1)
        item['layout_id']=f'recon-image-{number}'
        sheets[name].insert_image(rn,col,item['path'],{**options,'x_scale':scale,'y_scale':scale,'description':item['layout_id']})
        images.append({**item,'excel_sheet':name})
    wb.close()
    return images


def _write_qianchuan(sheet,dataset,formats):
    from .qianchuan import COLUMNS
    tables=dataset.table('qianchuan_monthly')
    for col,label in enumerate(COLUMNS):sheet.write_string(0,col,label)
    if dataset.context.get('qianchuan_access')=='permission_denied':
        sheet.write_comment(0,0,'千川无访问权限，金额未知；其他账户正常汇总，千川相关指标待补充。缺失金额不代表0。')
        sheet.write_string(2,16,'赠款余额(元)');sheet.write_string(2,17,'非赠款余额(元)')
        return
    q=dataset.context.get('qianchuan')
    rows=tables[0].rows if tables else [r for t in dataset.table('qianchuan') for r in t.rows]
    for rn in range(1,max(4,len(rows)+1)):
        if rn<=len(rows):
            for col,key in enumerate(COLUMNS):
                value=rows[rn-1].values.get(key)
                if col and isinstance(value,str) and value.strip() not in ['','-']:
                    try:value=decimal(value,blank_zero=False)
                    except DataError:pass
                _cell(sheet,rn,col,value,ColumnStub('string' if col==0 else 'decimal'),formats)
        if rn==2:
            sheet.write_string(2,16,'赠款余额(元)');sheet.write_string(2,17,'非赠款余额(元)')
        if rn==3:
            if q and q['grant'] is not None and q['cash'] is not None:
                sheet.write_number(3,16,float(q['grant']),formats['amount']);sheet.write_number(3,17,float(q['cash']),formats['amount'])
                sheet.write_comment(3,16,f"取自 {q['date']} 的赠款余额");sheet.write_comment(3,17,f"取自 {q['date']} 的非赠款余额")
            else:
                sheet.write_formula(3,16,'=NA()',formats['amount'],'#N/A');sheet.write_formula(3,17,'=NA()',formats['amount'],'#N/A')
                if q:
                    note='目标月无数据，未返回历史月末余额；无数据为正常状态。' if q.get('data_status')=='no_activity' else f"官方数据没有 {q['date']} 的余额记录；保留实际日期和金额，未用其他日期或0代替月末余额。"
                    sheet.write_comment(3,16,note);sheet.write_comment(3,17,note)
                    sheet.write_string(4,16,'当月无数据' if q.get('data_status')=='no_activity' else '月末无余额记录')
    sources=dataset.table('qianchuan')
    if len(sources)>1:
        sheet.write_comment(0,0,'主表为全部绑定千川账户的按日合计；每个账户的原始报表完整列于下方B:O列。Q4/R4为各账户月末余额合计。')
        rn=max(35,len(rows)+3)
        for table in sources:
            source=table.rows[0].source if table.rows else {}
            sheet.write_string(rn,0,f"广告账户：{source.get('entity_name') or ''}（{source.get('entity_id') or ''}）");rn+=1
            for col,label in enumerate(COLUMNS):sheet.write_string(rn,col+1,label)
            rn+=1
            for row in table.rows:
                for col,key in enumerate(COLUMNS):
                    value=row.values.get(key)
                    if col and isinstance(value,str) and value.strip() not in ['','-']:
                        try:value=decimal(value,blank_zero=False)
                        except DataError:pass
                    _cell(sheet,rn,col+1,value,ColumnStub('string' if col==0 else 'decimal'),formats)
                rn+=1
            rn+=2


def _wallet_balance_grid(dataset,start_col=8):
    images=[i for i in dataset.context.get('images',[]) if _image_sheet(i)=='管家账户']
    if not any(i.get('wallet_balance') for i in images):return {}
    # A retry can leave several captures of one subject. Use its latest capture
    # once, including a missing amount rather than silently using an old value.
    subjects={}
    for i,image in enumerate(images):
        key=image.get('entity_name') or image.get('entity_id') or f'capture-{i}'
        previous=subjects.get(key)
        if previous is None or (image.get('captured_at') or '') >= (previous.get('captured_at') or ''):subjects[key]=image
    rows=[]
    for image in subjects.values():
        evidence=image.get('wallet_balance') or {}
        unopened=evidence.get('status')=='not_opened'
        amount=None
        if not unopened and evidence.get('kind')=='current' and evidence.get('total') is not None:
            try:amount=decimal(evidence['total'],blank_zero=False)
            except DataError:pass
            if amount is not None and not amount.is_finite():amount=None
        captured=image.get('captured_at')
        if captured:
            captured=datetime.fromisoformat(captured.replace('Z','+00:00')).astimezone(timezone(timedelta(hours=8))).strftime('%Y-%m-%d %H:%M:%S')
        rows.append([image.get('entity_name') or image.get('entity_id') or '未标识主体',amount,
                     '未开通' if unopened else '已开通' if amount is not None else '余额待采集',captured or ''])
    amount_col=xlsxwriter.utility.xl_col_to_name(start_col+1)
    cells=f'{amount_col}11:{amount_col}{10+len(rows)}'
    total=sum((r[1] for r in rows if r[1] is not None),Decimal(0))
    complete=all(r[1] is not None or r[2]=='未开通' for r in rows)
    expected=sum(r[2]!='未开通' for r in rows)
    formula=f'=SUM({cells})' if complete else f'=IF(COUNT({cells})={expected},SUM({cells}),NA())'
    grid={0:{start_col:'管家账户余额汇总'},
          1:{start_col:'余额口径：截图采集时的总金额'},
          2:{start_col:'未开通主体不计入合计；非月末余额'},
          6:{start_col:'总余额（元）',start_col+1:(formula,total if complete else '#N/A')},
          9:{start_col+i:label for i,label in enumerate(['账户主体','总余额（元）','账户状态','截图时间（北京时间）'])}}
    for rn,row in enumerate(rows,10):grid[rn]={start_col+i:value for i,value in enumerate(row)}
    return grid


def _image_sheet(image):
    explicit=image.get('excel_sheet') or image.get('target_sheet')
    if explicit:return sheet_name(explicit)
    label=sheet_name(image.get('sheet') or '')
    if '千川' in label:return '千川账户'
    if '涉税' in label:return None  # Tax sources stay in the delivery artifacts when there is no report destination.
    if '资金账单' in label and '月汇总' in label:return '余额截图'
    if label in ['资金账单','结算账单','管家账户','保证金账户','余额截图']:return label
    if '保证金' in label:return '保证金账户'
    if '管家' in label:return '管家账户'
    if '余额' in label:return '余额截图'
    if '结算' in label:return '结算账单'
    if '资金' in label or '货款' in label:return '资金账单'
    return label


def _fund_month_balance(dataset):
    matches=[b for b in dataset.balances if b.get('account')=='货款账户'
             and b.get('period')==dataset.context['month'] and b.get('kind')=='monthly'
             and b.get('closing') is not None and b.get('source')]
    return matches[0] if len(matches)==1 else None


def _write_template_summary(sheet,dataset,formats,layout):
    grid={}
    # Preserve supplied shop codes and opening balances as source annotations.
    # They do not become verified monthly balance evidence in the rule engine.
    for coord,value in layout.cells('核对汇总').items():
        rn,col=xlsxwriter.utility.xl_cell_to_rowcol(coord)
        if col in [0,10,12,15] or rn==0:grid.setdefault(rn,{})[col]=value
    formulas=dataset.context['rule_snapshot']['summary_formulas']
    fees=next(t for t in dataset.tables if t.table_id=='fee_summary')
    for r in fees.rows:
        v=r.values;rn=v['模板行号']-1
        grid.setdefault(rn,{}).update({1:v['资金账户'],2:v['费用类型'],3:(formulas[f'D{rn+1}'],v['金额'])})
        if grid[rn].get(0) is None:grid[rn][0]=v.get('店铺编码')
    check=dataset.checks[0]
    for i,(label,value) in enumerate([('货款结算金额',check['计算值']),('结算账单金额',check['实际值']),('校验',check['差异'])],1):
        grid.setdefault(i,{})[6]=label;grid[i][7]=(formulas[f'H{i+1}'],value)
    balances=next(t for t in dataset.tables if t.table_id=='account_balances')
    for i,r in enumerate(balances.rows,1):
        v=r.values;grid.setdefault(i,{})[11]=v['资金账户']
        opening=v['期初余额']
        if opening is None and v['资金账户']!='千川账户':opening=grid[i].get(12)
        grid[i][12]=opening
        movement=sum((row.values['金额'] for row in fees.rows if row.values['资金账户']==v['资金账户'] and row.values['金额'] is not None),Decimal(0))
        if v['资金账户']=='千川账户':
            q=dataset.context.get('qianchuan')
            movement=q['consumption']+q['recharge'] if q else None
        closing=decimal(opening)+movement if opening is not None and movement is not None else None
        grid[i][13]=(formulas[f'N{i+1}'],movement);grid[i][14]=(formulas[f'O{i+1}'],closing)
        if v['资金账户']=='千川账户':
            q=dataset.context.get('qianchuan')
            grid[i][15]=(formulas['P6'],closing-q['grant']-q['cash'] if closing is not None and q and q['grant'] is not None and q['cash'] is not None else '#N/A')
            if dataset.context.get('qianchuan_access')=='permission_denied':
                grid[i][12]=None
                grid[i][15]=(formulas['P6'],'')
                sheet.write_comment(i,15,'千川无访问权限，相关数据留空；保留模板公式，未将缺失金额认定为0。')
            if q and (q['grant'] is None or q['cash'] is None):
                sheet.write_comment(i,15,'官方查询或报表未返回完整月末余额，正常采集已完成；#N/A表示该余额核对没有数值依据，不是权限或采集异常。')
        if v['差异'] is not None and v['资金账户']!='千川账户':
            grid[i][15]=('=+O2-余额截图!F27',v['差异']) if i==1 and _fund_month_balance(dataset) is not None else v['差异']
        elif v['资金账户']!='千川账户' and isinstance(grid[i].get(15),tuple):
            formula=grid[i][15][0]
            # Retain balance-check formulas, recaching only when the referenced
            # source cell is available. Formula evaluation in Excel remains on.
            ref=re.fullmatch(r"=\+?O[2-5]-'?([^'!]+)'?!([A-Z]+[0-9]+)",formula)
            reference=layout.cells(ref[1]).get(ref[2]) if ref and ref[1] in layout.sheets else None
            if isinstance(reference,tuple):reference=reference[1]
            grid[i][15]=(formula,closing-decimal(reference) if closing is not None and reference is not None else '#N/A')
        if v['期初余额'] is None:
            sheet.write_comment(i,12,'期初余额尚未核实；千川未知余额不补0，保留模板公式和缺失缓存。其他账户原表参考值不代表已核实余额，详见同批次result.json。')
    for rn in sorted(grid):
        for col,value in sorted(grid[rn].items()):
            if isinstance(value,tuple):
                cache=float(value[1]) if isinstance(value[1],(Decimal,int,float)) else value[1] if value[1] is not None else '#N/A'
                coord=xlsxwriter.utility.xl_rowcol_to_cell(rn,col)
                if dataset.context.get('qianchuan_access')=='permission_denied' and (col==3 and grid[rn].get(1)=='千川账户' or coord in ['N6','O6','P6']):
                    cache=''
                if coord in dataset.context['rule_snapshot'].get('array_formulas',[]):
                    sheet.write_array_formula(rn,col,rn,col,value[0],formats['amount'],cache)
                else:sheet.write_formula(rn,col,value[0],formats['amount'],cache)
            else:_cell(sheet,rn,col,value,ColumnStub('decimal' if col in [3,7,12,13,14,15] else 'string'),formats)


class ColumnStub:
    def __init__(self,kind):self.type=kind


def grid_page(db_path: Path,table_id: str,offset=0,limit=500):
    if offset<0 or not 1<=limit<=2000:raise ValueError('offset >= 0 and 1 <= limit <= 2000 required')
    with sqlite3.connect(f'file:{db_path.resolve()}?mode=ro',uri=True) as db:
        result=db.execute('SELECT schema_json FROM tables WHERE table_id=?',(table_id,)).fetchone()
        if not result:raise ValueError('unknown table_id')
        schema=json.loads(result[0]);columns=schema['columns']
        rows=[{'row_id':r[0],'values':[json.loads(r[1]).get(c['key']) for c in columns],'source':json.loads(r[2])}
              for r in db.execute('SELECT row_id,values_json,source_json FROM rows WHERE table_id=? ORDER BY ordinal LIMIT ? OFFSET ?',(table_id,limit,offset))]
    return {'schema_version':'1.0','table_id':table_id,'columns':columns,'offset':offset,'total':schema['row_count'],
            'next_offset':offset+len(rows) if offset+len(rows)<schema['row_count'] else None,'rows':rows}
