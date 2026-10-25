"""Exercise the real collectors for every shop and retain blocker evidence.

Does not upload, open accounts, or alter permissions. Each browser task belongs
to one shop. Checkpoints allow the audit to resume without repeating downloads.
"""
from __future__ import annotations

import argparse
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timedelta
import json
from pathlib import Path
import re
import sys
import threading
import time
import traceback
from urllib.parse import urlsplit

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src'))
from douyin_settle_recon.collect import UI, KINDS, URL, _monthly, _account
from douyin_settle_recon.collect_qianchuan import collect_all, ENTRY
from douyin_settle_recon.model import DataError, dump_json, month_bounds
from douyin_settle_recon.shopdesk import ShopDesk, ShopDeskError

PAGES = {
    'fund': ('资金账单', '资金→账单管理→资金账单→月汇总→生成报表→历史报表下载', URL),
    'settlement': ('结算账单', '资金→账单管理→结算账单→月汇总→下载明细→历史报表下载', 'https://fxg.jinritemai.com/ffa/fxg-bill/settlement-detail-bill'),
    'base_deposit': ('基础保证金', '资金→保证金账户→基础保证金→查询→导出明细', 'https://fxg.jinritemai.com/ffa/content/shop_deposit'),
    'experience_deposit': ('体验保证金', '资金→保证金账户→体验保证金→查询→导出明细', 'https://fxg.jinritemai.com/ffa/content/shop_deposit'),
    'steward': ('管家账户', '资金→账户中心→管家账户／抖店钱包→详情→切换主体→余额明细→生成报表→历史报表', 'https://fxg.jinritemai.com/ffa/fund-control/account-center'),
    'tax': ('涉税信息报送数据', '资金→账单管理→涉税信息报送数据→选择主体→按月度看→查询→下载明细→历史报表', 'https://fxg.jinritemai.com/ffa/fxg-bill/tax-bill'),
    'qianchuan': ('千川账户', '抖店后台→巨量千川（流量→千川推广）→我的资金→账户明细→流水对账→日结报表→选择月份→下载明细', ENTRY),
}
GOOD = {'downloaded', 'no_activity', 'not_opened'}
STATUS = {'permission': '权限不足', 'real_name': '账号实名认证受限', 'login': '登录／授权失效',
          'verification': '安全验证／平台限流', 'download': '导出／下载失败',
          'period_missing': '目标月份无可用账单',
          'qianchuan_not_opened': '千川入口待开户／绑定',
          'empty_table': '程序空流水识别异常',
          'validation': '下载数据校验失败', 'unconfirmed': '页面或程序异常，待复核',
          'downloaded': '已实际下载', 'no_activity': '已查询，当月无流水', 'not_opened': '账户未开通',
          'error': '抓取失败（原因见阻塞清单）', 'passed': '全部来源完成', 'blocked': '存在采集阻塞'}
PERMISSION = ('暂无权限', '无访问权限', '暂无访问权限', '没有访问权限', '没有权限访问',
              '没有访问该页面的权限', '没有该页面权限', '权限不足', '无权访问',
              '您没有权限', '你没有权限', '当前账号无权限', '无此权限', '未开通该权限',
              '权限被主账号调整无法访问', '权限点:「', '权限点：「')
PRINT_LOCK = threading.Lock()


def now():
    return datetime.now().astimezone().isoformat()


def emit(value):
    with PRINT_LOCK:
        print(json.dumps(value, ensure_ascii=False), flush=True)


def restriction(state):
    text = state.get('text', '')
    path = urlsplit(state.get('url', '')).path
    for code, words in [
        ('real_name', ('需要实名认证后访问', '你的账号未进行实名认证', '您的账号未进行实名认证', '请先完成实名认证')),
        ('permission', PERMISSION),
        ('verification', ('请完成安全验证', '请先通过安全验证', '访问过于频繁', '操作频繁，请稍后再试')),
    ]:
        for word in words:
            if word in text:
                lines = text.splitlines()
                i = next(i for i, line in enumerate(lines) if word in line)
                return code, '；'.join(lines[max(0, i - 1):i + 4])[:650]
    if re.search(r'/(?:account/)?login(?:/|$)', path, re.I) or '登录已失效' in text or ('手机号登录' in text and '密码登录' in text):
        return 'login', '当前业务页面要求登录或重新授权'
    return None


def classify(error, state):
    # Prefer observed official messages over the collector's generic timeout.
    found = restriction(state)
    if found:
        return found
    message = str(error)
    if any(x in message for x in PERMISSION):
        return 'permission', message
    if 'needs_real_name' in message:
        return 'real_name', message
    if 'needs_login' in message:
        return 'login', message
    if isinstance(error, DataError):
        return 'validation', message
    if 'no completed official download within timeout' in message:
        return 'download', '点击官方导出后，在等待时限内未收到完整下载文件'
    if 'official download was interrupted' in message:
        return 'download', '官方文件下载被取消或中断'
    if any(x in message.lower() for x in ['download', '下载', 'export', '报表生成失败']):
        return 'download', message
    return 'unconfirmed', message


def action_for(code, role, points=()):
    if code == 'permission':
        named = '；本页明确提示的权限项：' + '、'.join(points) if points else ''
        if role == 'qianchuan':
            return ('请主账号在抖店【店铺→子账号管理→岗位管理／账号管理→修改权限】检查当前登录账号的【流量→千川推广】权限；'
                    '若已进入千川后被拦截，请账户管理员检查对应千川广告账户的资金／账户明细／流水对账权限' + named)
        return ('请主账号在抖店【店铺→子账号管理→岗位管理／账号管理→修改权限】开通本行业务页面及查询、导出、历史报表下载权限；具体权限项以后台和本页提示为准' + named)
    if code == 'real_name':
        return '请当前登录账号持有人在【资金→账户中心→立即认证】按页面要求完成账号实名认证后复查；不能据此认定店铺经营主体未认证'
    if code == 'login':
        return '请在 ShopDesk 重新登录该店铺；千川另需确认绑定广告账户的登录／授权状态，恢复后重跑对应来源'
    if code == 'verification':
        return '请完成平台安全验证；如为访问频率限制，待限制解除后重试'
    if code == 'not_opened':
        return '官方页面明确显示账户尚未开通；按业务需要决定是否开通，无需据此修改子账号权限'
    if code == 'download':
        return '先复查对应来源的生成报表／历史报表下载页面；如页面无权限提示请管理员授权，否则需排查平台下载及程序适配'
    if code == 'validation':
        return '下载已执行，但内容或格式校验失败；需核对官方文件、月份、账号和程序格式适配，不能直接认定缺权限'
    if code == 'period_missing':
        return '请业务核实该店该月是否应有业务、是否需要纳入任务；如需要，请与平台核实该月份账单或余额的提供方式，程序也需单列处理目标月汇总缺失。已清退店铺需额外确认处理范围；当前证据不支持仅据此增加权限'
    if code == 'qianchuan_not_opened':
        return '请业务负责人确认本店是否需要千川账单；如已有广告账户，核实与当前店铺和登录账号的关联，如需新开户，由有授权的负责人在该页完成。当前入口展示开户表单，不能仅按子账号权限不足处理'
    if code == 'empty_table':
        return '需开发修正保证金表格的空流水／占位行判断，再按官方0条查询结果处理无流水分支；当前已有页面访问及查询能力，无需为该问题额外开通页面权限'
    return '需要复核页面加载、控件或主体切换及抓取程序；当前证据不足以要求客户开通权限'


class AuditUI(UI):
    """Stop promptly at explicit denials, retaining the real blocked page."""
    role = None

    def wait(self, fn, timeout=35):
        def checked():
            state = self.snapshot()
            path = urlsplit(state['url']).path
            # Ignore the previous application's DOM during an in-flight navigate.
            relevant = self.role and (
                path == urlsplit(PAGES[self.role][2]).path
                or self.role == 'qianchuan' and path.startswith('/ffa/ad/promotion-v2/')
                or self.role == 'qianchuan' and urlsplit(state['url']).hostname == 'qianchuan.jinritemai.com'
                or self.role == 'steward' and '/fund-control/' in path
                or path.endswith('/history-report')
            )
            denied = restriction(state) if relevant else None
            if denied:
                raise ShopDeskError(denied[1])
            return fn()
        return super().wait(checked, timeout)


def capture_issue(ui, root, shop, receipt, error=None, month='2026-09'):
    try:
        state = ui.snapshot()
    except ShopDeskError:
        state = {}
    observed_error = error or sys.exc_info()[1] or receipt.get('error', '')
    code, reason = classify(observed_error, state)
    tb = error.__traceback__ if isinstance(error, BaseException) else sys.exc_info()[2]
    frames = traceback.extract_tb(tb) if tb else []
    stage = ''
    names = {frame.name for frame in frames}
    if '_select_entity' in names:
        stage = '主体选择／切换条件未能确认'
    elif '_select_tax_month' in names or '_tax_control' in names:
        stage = '涉税月份选择控件未能按预期完成'
    elif any('queried' in (frame.line or '') for frame in frames):
        stage = '官方查询回执未能在时限内确认'
    elif '_report_ready' in names:
        stage = '官方报表任务未能在时限内确认已生成'
    elif '_generate_month_report' in names:
        stage = '官方报表生成流程未能完成'
    elif any('month_row' in (frame.line or '') for frame in frames):
        stage = '目标月份的数据行未能在页面确认'
    elif any('_accounts' in (frame.line or '') for frame in frames):
        stage = '千川入口可访问，但切换后的绑定广告账户列表未能确认'
    elif '_select_month' in names:
        stage = '千川日结页面的目标月份选择未能完成'
    elif '_query' in names:
        stage = '千川官方日结查询回执未能在时限内确认'
    elif '_open' in names:
        stage = '千川账户明细或财务嵌入页面未能按预期加载'
    elif '_entry' in names:
        stage = '千川入口应用或绑定账户选择页面未能按预期加载'
    if code == 'unconfirmed' and 'condition did not become ready' in reason:
        reason = (stage or '页面加载或操作后的验证条件未能在时限内确认') + '；当前未观察到明确的权限拦截提示，需复核抓取程序及页面状态'
    item = {'role': receipt['role'], 'entity': receipt.get('entity'),
            'advertiser_id': receipt.get('advertiser_id'), 'status': code,
            'reason': reason, 'collector_error': str(error or receipt.get('error', ''))[:1200],
            'checked_at': now(), 'actual_url': '', 'page_excerpt': state.get('text', '')[-6000:]}
    item['stage'] = stage
    item['exception_type'] = type(observed_error).__name__
    item['exception_frames'] = [{'function': f.name, 'file': Path(f.filename).name, 'line': f.lineno,
                                  'source': f.line} for f in frames]
    if receipt['role'] == 'tax' and code == 'unconfirmed' and stage == '官方查询回执未能在时限内确认':
        cursor = tb
        while cursor:
            if cursor.tb_frame.f_code.co_name == '_tax_export':
                query = cursor.tb_frame.f_locals.get('request', {})
                scope = query.get('body', {})
                item['query_diagnostics'] = {'request_observed': bool(query),
                    'scope': {k: scope.get(k) for k in ['member_id', 'date_type', 'bill_date', 'query_type']}}
                if not query:
                    item['reason'] = '正式抓取点击“查询”后，未观测到预期涉税查询请求，无法取得本次月份／主体查询回执，等待20秒超时；现场页面可访问，需复核查询控件与请求识别逻辑'
                elif query.get('id'):
                    try:
                        raw = ui.bridge.command(ui.eid, 'Network.getResponseBody', {'requestId': query['id']})
                        response = json.loads(raw['result']['body'])
                        item['query_diagnostics']['response'] = {k: response.get(k) for k in ['code', 'msg', 'message', 'total']}
                        item['query_diagnostics']['has_data'] = bool(response.get('data'))
                    except (ShopDeskError, KeyError, ValueError) as exc:
                        item['query_diagnostics']['response_read_error'] = str(exc)[:500]
                        item['reason'] = '已观测到涉税查询请求，但正式抓取未能在20秒内取得可读取的官方响应；页面访问未见权限拦截，需复核响应读取或平台返回情况'
                break
            cursor = cursor.tb_next
    if state.get('url'):
        parsed = urlsplit(state['url'])
        item['actual_url'] = parsed.scheme + '://' + parsed.netloc + parsed.path
    item['permission_points'] = re.findall(r'权限点[:：]\s*[「“]([^」”]+)[」”]', state.get('text', ''))
    item['action'] = action_for(code, item['role'], item['permission_points'])
    try:
        shot = ui.capture(root, shop['name'] + '-' + item['role'] + '-阻塞', month, width=1440, height=1000)
        item['screenshot'] = shot['path']
    except (ShopDeskError, OSError) as exc:
        item['screenshot_error'] = str(exc)
    return item


class Receipts(list):
    def __init__(self, values, on_error):
        super().__init__(values)
        self.on_error = on_error

    def append(self, value):
        super().append(value)
        if value.get('status') == 'error':
            self.on_error(value)


def audit_shop(root, month, shop, index, total, production_waits=False):
    directory = root / 'shops' / shop['id']
    directory.mkdir(parents=True, exist_ok=True)
    file = directory / 'audit.json'
    report = json.loads(file.read_text()) if file.exists() else {
        'shop': shop, 'month': month, 'started_at': now(), 'roles': [], 'issues': [], 'complete': False}
    if report.get('complete'):
        return report
    manifest_path = directory / 'collection.json'
    manifest = json.loads(manifest_path.read_text()) if manifest_path.exists() else {
        'schema_version': '1.0', 'shop': {'shopdesk_shop_uuid': shop['id'], 'platform_shop_id': None},
        'month': month, 'requested_kinds': list(KINDS), 'inputs': [], 'screenshots': [],
        'balances': [], 'collection': []}
    bridge = ShopDesk()
    eid = None
    try:
        eid = bridge.start(shop['id'], URL, wait_timeout=600)
        ui = (UI if production_waits else AuditUI)(bridge, eid)
        report['last_attempt_wait_mode'] = 'production' if production_waits else 'permission-aware audit'
        dump_json(directory / 'session.json', {'exploration_id': eid, 'shop_id': shop['id']})
        bridge.ready(eid)
        def on_error(receipt):
            report['issues'].append(capture_issue(ui, directory, shop, receipt, month=month))
            dump_json(file, report)
        manifest['collection'] = Receipts(manifest['collection'], on_error)
        # Each collector navigates and verifies its own destination. Even when
        # account center is denied, other independent sources must still run.
        for role in KINDS:
            if any(x['role'] == role for x in report['roles']):
                continue
            ui.role = role
            before = len(manifest['collection'])
            started = now()
            try:
                if role in ('fund', 'settlement'):
                    _monthly(ui, role, month, directory, manifest)
                elif role == 'qianchuan':
                    collect_all(ui, month, directory, manifest)
                else:
                    _account(ui, role, month, directory, manifest)
            except Exception as exc:
                # Keep a shop/role failure from skipping all remaining shops.
                receipt = {'role': role, 'status': 'error', 'error': f'{type(exc).__name__}: {exc}'}
                list.append(manifest['collection'], receipt)
                report['issues'].append(capture_issue(ui, directory, shop, receipt, exc, month))
            entries = manifest['collection'][before:]
            if not entries:
                # A resumed multi-subject collector may legitimately skip all
                # already downloaded subjects after enumerating them again.
                entries = [x for x in manifest['collection'] if x['role'] == role and x['status'] in GOOD]
            if not entries:
                receipt = {'role': role, 'status': 'error', 'error': 'Collector finished without a result receipt'}
                manifest['collection'].append(receipt)
                entries = [receipt]
            entry = {'role': role, 'started_at': started, 'finished_at': now(),
                     'complete': all(x['status'] in GOOD for x in entries),
                     'receipts': list(entries),
                     'downloaded_files': [x for x in manifest['inputs'] if x['role'] == role]}
            report['roles'].append(entry)
            dump_json(manifest_path, manifest)
            dump_json(file, report)
            emit({'index': index, 'total': total, 'shop': shop['name'], 'role': role,
                  'results': [x['status'] for x in entries], 'issues': len(report['issues'])})
        report['complete'] = len(report['roles']) == len(KINDS)
    except Exception as exc:
        report['session_error'] = str(exc)
        if getattr(exc, 'code', None) in {'EXPLORATION_WAIT_TIMEOUT', 'EXPLORATION_BUSY'}:
            # A shared local browser pool is not a shop permission failure.
            # Leave the shop pending so a later audit invocation retries it.
            report['complete'] = False
            return report
        # A browser/session failure affects all remaining untested sources.
        code = 'login' if 'needs_login' in str(exc) else 'unconfirmed'
        for role in KINDS:
            if any(x['role'] == role for x in report['roles']):
                continue
            issue = {'role': role, 'status': code, 'reason': str(exc), 'checked_at': now(),
                     'collector_error': str(exc), 'screenshot_error': '工作台未能建立可访问业务页面的会话，无法截取该页面',
                     'action': action_for(code, role)}
            report['issues'].append(issue)
            report['roles'].append({'role': role, 'complete': False,
                                    'receipts': [{'role': role, 'status': 'error', 'error': str(exc)}], 'downloaded_files': []})
        report['complete'] = len(report['roles']) == len(KINDS)
    finally:
        if eid:
            try:
                bridge.close(eid)
            except ShopDeskError as exc:
                report['close_error'] = str(exc)
        report['finished_at'] = now()
        dump_json(manifest_path, manifest)
        dump_json(file, report)
    return report


def prepare_transient_retries(root, shops):
    """Recheck failures that are not explicit permission/identity denials."""
    for shop in shops:
        directory = root / 'shops' / shop['id']
        file = directory / 'audit.json'
        if not file.is_file():
            continue
        report = json.loads(file.read_text())
        roles = {i['role'] for i in report['issues'] if i['status'] in {'unconfirmed', 'download'}}
        if not roles:
            continue
        history = directory / 'attempt-history'
        history.mkdir(exist_ok=True)
        number = len(list(history.glob('*.json'))) + 1
        dump_json(history / f'{number}.json', report)
        report['retry_history'] = [str(p) for p in sorted(history.glob('*.json'))]
        report['roles'] = [x for x in report['roles'] if x['role'] not in roles]
        report['issues'] = [x for x in report['issues'] if x['role'] not in roles]
        report['complete'] = False
        manifest_path = directory / 'collection.json'
        manifest = json.loads(manifest_path.read_text())
        manifest['collection'] = [x for x in manifest['collection'] if x['role'] not in roles or x['status'] in GOOD]
        dump_json(manifest_path, manifest)
        dump_json(file, report)
        emit({'retrying_transient': shop['name'], 'roles': sorted(roles)})


def explain_incomplete_dates(root, reports, month):
    """Explain the actual downloaded dates without assuming why days are absent."""
    from openpyxl import load_workbook
    start, end = month_bounds(month)
    expected = {(start + timedelta(days=i)).strftime('%Y-%m-%d') for i in range((end - start).days)}
    for report in reports:
        directory = root / 'shops' / report['shop']['id']
        manifest = json.loads((directory / 'collection.json').read_text())
        changed = False
        for issue in report['issues']:
            page_text = issue.get('page_excerpt', '')
            observed = restriction({'text': page_text, 'url': issue.get('actual_url', '')})
            if observed and issue['status'] != observed[0]:
                issue['classification_note'] = '根据已保存的官方页面原文复核分类，原程序错误原文仍保留'
                issue['status'], issue['reason'] = observed
                issue['action'] = action_for(issue['status'], issue['role'], issue.get('permission_points', []))
                changed = True
            empty_proof = directory / 'empty-deposit-proof.json'
            if issue['role'] == 'base_deposit' and issue['status'] in {'download', 'empty_table'} and empty_proof.is_file():
                proof = json.loads(empty_proof.read_text())
                if proof.get('rows') == [['暂无数据']] and re.search(r'共\s*0\s*条', proof.get('text', '')):
                    issue['status'] = 'empty_table'
                    issue['reason'] = '基础保证金官方页面显示“共0条／暂无数据”，实测表格读取返回占位行[["暂无数据"]]。当前程序把该占位行当作真实流水，仍触发导出，等待120秒未收到文件；属于程序空流水识别问题'
                    issue['action'] = action_for('empty_table', issue['role'])
                    issue['empty_table_proof'] = str(empty_proof)
                    issue.setdefault('error_screenshot', issue.get('screenshot'))
                    issue['screenshot'] = proof['screenshot']
                    changed = True
            if issue['role'] == 'qianchuan' and not observed:
                if '/account/register' in issue.get('actual_url', '') and '立即免费开户' in page_text and '免费开通巨量千川账户' in page_text:
                    issue.update(status='qianchuan_not_opened',
                                 reason='千川入口实际展示“免费开通巨量千川账户／立即免费开户”表单，未进入已绑定广告账户的资金页面；当前无法查询或下载本店千川日结数据',
                                 action=action_for('qianchuan_not_opened', 'qianchuan'),
                                 required_page_path='抖店后台→流量→千川推广→免费开户／核实已有账户关联')
                    changed = True
                elif issue['status'] == 'unconfirmed' and '收起导航' in page_text and page_text.split('收起导航', 1)[1].strip() in {'', '待办'}:
                    issue['reason'] = '千川入口已打开，但等待35秒后业务区域仍为空白，仅显示抖店外层导航，无法读取绑定账户；页面未出现明确的权限拦截提示'
                    issue['action'] = '需开发／平台复核千川前端加载、登录关联和页面返回；当前没有足够证据要求客户增加权限'
                    match = re.search(r'(?:^|\n)(?:发布|快速发布)\n([^\n]+)\n(?:企业店|个体店|旗舰店|专卖店|专营店|个人店)\n', page_text)
                    if match and match[1] != report['shop']['name']:
                        issue['observed_shop_name'] = match[1]
                        issue['reason'] += f'；工作台记录名为“{report["shop"]["name"]}”，页面店铺名为“{match[1]}”'
                        issue['action'] += '；另请业务核对该店是否更名，或工作台标签与当前登录店铺的映射是否正确'
                    changed = True
            if issue['role'] in {'fund', 'settlement'} and not observed:
                months = sorted(set(re.findall(r'(?m)^(20\d{2}-\d{2})(?=\s|$)', page_text)))
                closed = any(line.strip() == '已清退' for line in page_text.splitlines())
                totals = re.findall(r'共\s*(\d+)\s*条', page_text)
                full_list = bool(totals) and int(totals[-1]) == len(months)
                if months and month not in months and (closed or full_list):
                    issue['status'] = 'period_missing'
                    prefix = '页面显示店铺“已清退”；' if closed else f'业务页面可访问，已显示全部{len(months)}个月汇总条目；'
                    issue['reason'] = (prefix + f'当前可见月汇总最新月份为{max(months)}，未找到目标月份{month}的账单行。'
                                       '当前抓取链路等待该月行超时，尚不能据此确认目标月余额为0或认定缺权限。')
                    issue['visible_statement_months'] = months
                    issue['action'] = action_for('period_missing', issue['role'])
                    changed = True
            error_text = issue.get('collector_error', '')
            if issue['role'] != 'qianchuan' or not any(t in error_text for t in ['日结报表日期不完整', '缺少月末最后一天']):
                continue
            files = [x for x in manifest['inputs'] if x['role'] == 'qianchuan' and x.get('advertiser_id') == issue.get('advertiser_id')]
            actual = set()
            for item in files:
                with Path(item['path']).open('rb') as stream:
                    book = load_workbook(stream, read_only=True, data_only=True)
                    for sheet in book:
                        for row in sheet.iter_rows(values_only=True):
                            if row and re.fullmatch(r'\d{4}-\d{2}-\d{2}', str(row[0])):
                                actual.add(str(row[0]))
                    book.close()
            if not actual:
                continue
            missing = sorted(expected - actual)
            official_counts = [x['query']['total'] for x in manifest.get('screenshots', [])
                               if x.get('advertiser_id') == issue.get('advertiser_id') and 'total' in x.get('query', {})]
            full_shots = [x for x in manifest.get('screenshots', [])
                          if x.get('advertiser_id') == issue.get('advertiser_id') and Path(x.get('path', '')).is_file()]
            if full_shots:
                issue.setdefault('error_screenshot', issue.get('screenshot'))
                issue['screenshot'] = full_shots[-1]['path']
                issue['evidence_note'] = '采用官方日结查询完成时的完整明细截图；当前阻塞发生在其后的文件日期校验'
            issue['status'] = 'validation'
            issue['date_validation'] = {'expected_days': len(expected), 'actual_dates': sorted(actual), 'missing_dates': missing,
                                        'unexpected_dates': sorted(actual - expected), 'files': files, 'official_query_totals': official_counts}
            issue['reason'] = (f'官方千川日结文件已下载；当前程序对{month}期望{len(expected)}个日期，文件实际包含{len(actual)}个日期'
                               f'（{min(actual)}至{max(actual)}），缺少：' + '、'.join(missing) +
                               '。当前程序要求整月逐日记录，因此在下载后的日期完整性校验受阻。')
            if '缺少月末最后一天' in error_text:
                issue['reason'] = (f'官方千川日结文件已下载，仅包含{len(actual)}个日期（{min(actual)}至{max(actual)}），'
                                   f'缺少月末{end - timedelta(days=1):%Y-%m-%d}，当前无法按指定最后一天提取赠款及非赠款余额。')
            if official_counts:
                issue['reason'] += '官方同月查询返回记录数：' + '、'.join(map(str, sorted(set(official_counts)))) + '；查询结果与下载记录数可对照复查。'
            issue['action'] = '请开发核对官方日结导出对无流水日期的呈现规则及程序整月校验；保留原始文件并核对月末余额，当前证据不支持要求客户增加权限或推断缺失日期的金额'
            changed = True
        if changed:
            dump_json(directory / 'audit.json', report)


def run(root, month, workers=2, limit=None, retry_transient=False):
    month_bounds(month)
    root.mkdir(parents=True, exist_ok=True)
    listing = root / 'shops.json'
    if listing.exists():
        scope = json.loads(listing.read_text())
        if scope['month'] != month:
            raise ValueError('Audit month differs from the saved scope')
        shops = scope['shops']
    else:
        shops = ShopDesk().shops()
        dump_json(listing, {'checked_at': now(), 'month': month, 'shops': shops,
                            'method': 'real production collectors, official queries and downloads, no cloud upload'})
    if not shops:
        raise ValueError('No Douyin shops available')
    selected = shops[:limit] if limit else shops
    if retry_transient:
        prepare_transient_retries(root, selected)
    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(audit_shop, root, month, shop, i, len(shops), retry_transient) for i, shop in enumerate(selected, 1)]
        for future in as_completed(futures):
            report = future.result()
            emit({'shop_done': report['shop']['name'], 'issues': len(report['issues']), 'complete': report['complete']})
    reports = []
    for shop in shops:
        file = root / 'shops' / shop['id'] / 'audit.json'
        if file.exists():
            reports.append(json.loads(file.read_text()))
    explain_incomplete_dates(root, reports, month)
    dump_json(root / 'progress.json', {'checked_at': now(), 'scope': len(shops),
              'completed': sum(r['complete'] for r in reports),
              'blocked_shops': sum(bool(r['issues']) for r in reports)})


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--month', default='2026-09')
    parser.add_argument('--workers', type=int, choices=(1, 2), default=2)
    parser.add_argument('--limit', type=int)
    parser.add_argument('--retry-transient', action='store_true')
    args = parser.parse_args()
    run(args.output.resolve(), args.month, args.workers, args.limit, args.retry_transient)
