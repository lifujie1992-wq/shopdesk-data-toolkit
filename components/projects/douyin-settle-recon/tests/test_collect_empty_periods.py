from recon_samples import q_receipt
from copy import deepcopy
import json
from pathlib import Path
from types import SimpleNamespace

from openpyxl import load_workbook
import pytest

from douyin_settle_recon import collect
from douyin_settle_recon.monthly import _complete_collection
from douyin_settle_recon.shopdesk import ShopDeskError
from test_local_pipeline import TEMPLATE


class MonthlyUI:
    eid='owned'
    bridge=SimpleNamespace(command=lambda *a:None,ready=lambda *a:None)

    def __init__(self,pages):self.pages=deepcopy(pages);self.page=0;self.downloads=0
    def wait(self,fn,**kwargs):
        for _ in range(3):
            value=fn()
            if value:return value
        raise ShopDeskError('not ready')
    def month_summary(self):return self.pages[self.page]
    def next_month_page(self):self.page+=1
    def snapshot(self):return {'url':collect.URL,'text':'资金流水明细\n结算明细\n月汇总'}
    def click(self,label):assert label=='月汇总'
    def capture(self,*a):return {'path':'empty.png','sheet':'月汇总'}
    def download(self,*a):raise AssertionError('Empty month must not export')


def page(rows,total,**kwargs):return {'rows':rows,'total':total,'loading':False,'next':False,'empty':not rows,**kwargs}


@pytest.mark.parametrize('kind',['fund','settlement'])
@pytest.mark.parametrize('rows',[[],[['2026-06','1']],[['2026-06','1'],['2026-05','1']]])
def test_month_absent_from_complete_list_is_successful_without_export(tmp_path,kind,rows):
    ui=MonthlyUI([page(rows,len(rows))]);manifest={'screenshots':[],'balances':[],'collection':[]}
    collect._monthly(ui,kind,'2026-09',tmp_path,manifest)
    assert manifest['collection'][0]['status']=='no_activity'
    assert manifest['collection'][0]['query']['complete_list']
    assert manifest['balances']==[] and len(manifest['screenshots'])==1


def test_month_summary_checks_next_page_before_declaring_no_data():
    target=['2026-09','0','0','0','0','0']
    ui=MonthlyUI([page([['2026-10','1']],2,next=True),page([target],2)])
    assert collect._month_summary_result(ui,'2026-09')['row']==target
    ui=MonthlyUI([page([['2026-10','1']],2,next=True),page([['2026-08','1']],2)])
    result=collect._month_summary_result(ui,'2026-09')
    assert result['row'] is None and result['visible_months']==['2026-08','2026-10']


@pytest.mark.parametrize('state',[page([],None),page([['2026-06','1']],25),page([],0,loading=True)])
def test_loading_unknown_or_incomplete_list_is_never_no_data(state):
    with pytest.raises(ShopDeskError):collect._month_summary_result(MonthlyUI([state]),'2026-09')


def test_zero_settlement_row_skips_generation_and_download(tmp_path):
    ui=MonthlyUI([page([['2026-09','¥0.00','0','¥0.00','0','¥0.00']],1)])
    manifest={'screenshots':[],'balances':[],'collection':[]}
    collect._monthly(ui,'settlement','2026-09',tmp_path,manifest)
    assert manifest['collection'][0]['status']=='no_activity'


@pytest.mark.parametrize('kind',['base_deposit','experience_deposit'])
def test_deposit_placeholder_is_a_normal_empty_ledger(tmp_path,kind):
    class Empty(MonthlyUI):
        def snapshot(self):return {'text':'交易流水\n共 0 条\n暂无数据'}
        def rows(self):return [['暂无数据']]
        def js(self,script,*args):return ['',''] if 'map(e=>e.value)' in script else True
    ui=Empty([]);manifest={'screenshots':[],'inputs':[],'collection':[]}
    collect._account(ui,kind,'2026-09',tmp_path,manifest)
    assert manifest['collection'][0]['status']=='no_activity'
    assert manifest['inputs']==[]


def test_entire_shop_empty_still_builds_and_completes_collection(tmp_path,monkeypatch):
    bridge=SimpleNamespace(shops=lambda:[{'id':'shop','name':'空月份店铺'}],start=lambda *a:'owned',
                           ready=lambda *a:None,command=lambda *a:None,close=lambda *a:None)
    monkeypatch.setattr(collect,'ShopDesk',lambda:bridge)
    names=iter(['平台初始店名','平台改名后的店名'])
    monkeypatch.setattr(collect,'read_shop_identity',lambda ui:{'platform':'douyin','shop_id':'123','shop_name':next(names)})
    def empty(ui,kind,month,root,manifest):manifest['collection'].append(q_receipt(month=month) if kind=='qianchuan' else {'role':kind,'status':'no_activity'})
    monkeypatch.setattr(collect,'_monthly',empty);monkeypatch.setattr(collect,'_account',empty)
    monkeypatch.setattr('douyin_settle_recon.collect_qianchuan.collect_all',
                        lambda ui,month,root,manifest:empty(ui,'qianchuan',month,root,manifest))
    result=collect.collect_and_build(kinds=list(collect.KINDS),month='2026-09',shopdesk_shop_uuid='shop',
                                    platform_shop_id='123',output=tmp_path,template=TEMPLATE)
    assert _complete_collection(result['collection_manifest'],list(collect.KINDS))
    assert result['issues']==0 and result['conclusion']=='无数据'
    assert Path(result['report']).name=='202609【抖音】Y-平台改名后的店名.xlsx'
    collection=json.loads(Path(result['collection_manifest']).read_text())
    assert collection['shop']['shop_name']=='平台改名后的店名' and collection['shop']['platform_shop_id']=='123'
    data=json.loads((Path(result['run_dir'])/'result.json').read_text())
    assert data['issues']==[]
    book=load_workbook(result['report'],data_only=True)
    assert book['核对汇总']['D44'].value==0 and book['核对汇总']['D45'].value==0
    assert book['千川账户']['Q4'].value=='#N/A' and book['千川账户']['Q5'].value=='当月无数据'
    book.close()


def test_header_only_qianchuan_without_empty_proof_is_failed(tmp_path):
    from openpyxl import Workbook
    from douyin_settle_recon.pipeline import build
    from douyin_settle_recon.qianchuan import COLUMNS
    from douyin_settle_recon.schema import FUND_COLUMNS
    source=tmp_path/'empty.xlsx';wb=Workbook();ws=wb.active;ws.title='资金账单';ws.append(FUND_COLUMNS)
    ws=wb.create_sheet('结算账单');ws.append(['订单号','结算时间','商品名称','结算金额'])
    ws=wb.create_sheet('千川账户');ws.append(COLUMNS)
    wb.save(source);wb.close()
    receipts=[{'role':k,'status':'downloaded' if k in ['fund','settlement','qianchuan'] else 'no_activity'} for k in collect.KINDS]
    result=build([source],month='2026-09',shop_name='空文件',output=tmp_path/'runs',template=TEMPLATE,
                 collection_context={'collection':receipts})
    report=json.loads((Path(result['run_dir'])/'result.json').read_text())
    assert report['context']['qianchuan_result']['status']=='failed'
    assert not Path(result['report']).with_suffix('.xlsx.gz').exists()


@pytest.mark.parametrize('role,message,status',[
    ('qianchuan','千川入口权限不足','permission_denied'),
    ('qianchuan','暂无权限','permission_denied'),
    ('qianchuan','finance UI condition did not become ready','error'),
    ('qianchuan','下载明细条数不一致','error'),
    ('fund','权限不足','error'),
])
def test_only_explicit_qianchuan_denial_permits_partial_delivery(tmp_path,role,message,status):
    from douyin_settle_recon.collection_status import collection_error
    receipt=collection_error(role,ShopDeskError(message))
    assert receipt['status']==status
    p=tmp_path/'collection.json';p.write_text(json.dumps({'collection':[receipt]}))
    assert _complete_collection(p,[role])==(status=='permission_denied')


def test_empty_shop_with_qianchuan_denied_is_still_exported(tmp_path,monkeypatch):
    bridge=SimpleNamespace(shops=lambda:[{'id':'shop','name':'空月份店铺'}],start=lambda *a:'owned',
                           ready=lambda *a:None,command=lambda *a:None,close=lambda *a:None)
    monkeypatch.setattr(collect,'ShopDesk',lambda:bridge)
    monkeypatch.setattr(collect,'read_shop_identity',lambda ui:{'platform':'douyin','shop_id':'123','shop_name':'空月份店铺'})
    def empty(ui,kind,month,root,manifest):manifest['collection'].append(q_receipt(month=month) if kind=='qianchuan' else {'role':kind,'status':'no_activity'})
    def denied(*args,**kwargs):raise ShopDeskError('千川入口权限不足')
    monkeypatch.setattr(collect,'_monthly',empty);monkeypatch.setattr(collect,'_account',empty)
    monkeypatch.setattr('douyin_settle_recon.collect_qianchuan.collect_all',denied)
    result=collect.collect_and_build(kinds=list(collect.KINDS),month='2026-09',shopdesk_shop_uuid='shop',output=tmp_path,template=TEMPLATE)
    assert _complete_collection(result['collection_manifest'],list(collect.KINDS))
    assert Path(result['report']).with_suffix('.xlsx.gz').exists()
    assert result['partial_accounts']==['千川账户'] and result['upload_ready']
    manifest=json.loads(Path(result['collection_manifest']).read_text())
    assert manifest['collection'][-1]['status']=='permission_denied'
    assert result['conclusion']=='待确认'
    book=load_workbook(result['report'],data_only=True)
    assert book['千川账户']['A2'].value is None and book['千川账户']['Q4'].value is None
    for coord in ['D44','D45','M6','N6','O6','P6']:assert book['核对汇总'][coord].value is None,coord
    book.close()
    old_bytes=Path(result['report']).read_bytes()
    def unexpected(*args,**kwargs):raise AssertionError('successful sources must be reused')
    monkeypatch.setattr(collect,'_monthly',unexpected);monkeypatch.setattr(collect,'_account',unexpected)
    monkeypatch.setattr('douyin_settle_recon.collect_qianchuan.collect_all',
                        lambda ui,month,root,manifest:manifest['collection'].append(q_receipt(month=month)))
    complete=collect.collect_and_build(kinds=list(collect.KINDS),resume=Path(result['collection_manifest']),
        month='2026-09',shopdesk_shop_uuid='shop',output=tmp_path,template=TEMPLATE)
    assert complete['run_id']!=result['run_id'] and not complete['partial_accounts']
    assert Path(result['report']).read_bytes()==old_bytes
    book=load_workbook(complete['report'],data_only=True)
    assert book['核对汇总']['D44'].value==book['核对汇总']['D45'].value==0
    book.close()
