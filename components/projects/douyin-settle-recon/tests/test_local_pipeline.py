import csv
from decimal import Decimal
import gzip
import json
from pathlib import Path
import sqlite3

import openpyxl
import pytest

from douyin_settle_recon.model import DataError, Dataset, Row, Table, Column, decimal
from douyin_settle_recon.rules import TemplateRules, process
from douyin_settle_recon.schema import FUND_COLUMNS
from douyin_settle_recon.pipeline import build
from douyin_settle_recon.export import grid_page
from recon_samples import empty_collection

ROOT=Path(__file__).resolve().parents[1]
TEMPLATE=ROOT/'doc/【抖音】月度账单处理模版.xlsx'


def fund(**values):
    record={k:None for k in FUND_COLUMNS}
    record.update({'动账时间':'2026-08-31 23:59:59','动帐流水号':"'1234567890123456789012345",'动账方向':'入账','动账金额':Decimal('100.10'),'动账场景':'货款结算入账','订单实付应结':Decimal('101.10'),'运费实付':Decimal('1.00'),'平台服务费':Decimal('-1.00')})
    record.update(values)
    return record


def dataset(records):
    d=Dataset({'month':'2026-08','shop_name':'测试店','run_id':'test'})
    t=Table('fund','资金账单',[Column(k,k) for k in FUND_COLUMNS],role='fund')
    t.rows=[Row(r,{'artifact_id':'source','sheet':'fund','row':i}) for i,r in enumerate(records,2)]
    d.tables.append(t)
    return d


def test_template_precedence_and_negative_deposit():
    r=TemplateRules(TEMPLATE)
    assert r.classify({'备注':'退换货运费险，超时消费者赔付','动账场景':'货款结算入账'})[0]=='运费赔付'
    assert r.classify({'备注':'超时消费者赔付','动账场景':'消费者赔付'})[0]=='超时赔付'
    assert r.classify({'备注':'abc其他','动账场景':''})[0]=='abc'
    d=dataset([fund()])
    t=Table('deposit','体验保证金账户',[Column(k,k) for k in ['操作单号','操作时间','操作类型','操作金额','备注']],role='experience_deposit')
    t.rows=[Row({'操作单号':'888','操作时间':'2026-08-02','操作类型':'售后退款','操作金额':'-9.80','备注':'退款'},{'artifact_id':'deposit','row':2})]
    d.tables.append(t);p=process(d,r)
    assert p.rows[0].values['订单收入']==Decimal('101.10')
    assert p.rows[1].values['结算金额']==Decimal('-9.80')
    assert p.rows[1].values['资金账户']=='体验保证金账户'


def test_month_uses_accounting_date_not_order_date():
    d=dataset([fund(下单时间='2026-07-01'),fund(动账时间='2026-09-01 00:00:00')])
    p=process(d,TemplateRules(TEMPLATE))
    assert len(p.rows)==1
    assert any(i['code']=='outside_month' for i in d.issues)


def test_unknown_scene_does_not_turn_into_zero_or_success():
    d=dataset([fund(动账场景='未曾出现的新费用',备注=None)])
    p=process(d,TemplateRules(TEMPLATE))
    assert p.rows[0].values['费用类型'] is None
    assert d.context['conclusion']=='待确认'
    assert any(i['code']=='unmapped_scene' for i in d.issues)


def test_decimal_and_invalid_amount():
    assert decimal('0.10')+decimal('0.20')==Decimal('0.30')
    for value in ['N/A','NaN','Infinity',True]:
        with pytest.raises(DataError):decimal(value)


def test_official_deposit_currency_and_history_preserve_raw(tmp_path):
    from douyin_settle_recon.parse import import_file
    source=tmp_path/'deposit.csv'
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['操作单号','操作时间','操作类型','操作金额','备注','交易状态'])
        w.writerows([['\t7690489990928728381','\t2026-08-28 16:07:52','\t手动充值','\t+¥1,100.00','充值保证金','交易成功'],
                     ['old','2025-06-19','手动充值','+¥5000.00','旧数据','交易成功'],
                     ['refund','2026-08-29','售后退款','-￥9.80','退款','交易成功']])
    d=Dataset({'month':'2026-08','shop_name':'测试店'})
    import_file(d,source,tmp_path/'run',role='experience_deposit',account='体验保证金账户')
    p=process(d,TemplateRules(TEMPLATE))
    assert len(p.rows)==2
    assert [r.values['结算金额'] for r in p.rows]==[Decimal('1100.00'),Decimal('-9.80')]
    assert d.tables[0].rows[0].values['操作金额']=='\t+¥1,100.00'
    assert not any(i['code']=='invalid_amount' for i in d.issues)


def test_official_settlement_legend_is_retained_but_not_a_transaction(tmp_path):
    from douyin_settle_recon.parse import import_file,target_rows
    source=tmp_path/'settlement.csv'
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['结算时间','结算金额','订单号'])
        w.writerows([['','收入+支出',''],['2026-08-01','10.00','123']])
    d=Dataset({'month':'2026-08','shop_name':'测试店'})
    import_file(d,source,tmp_path/'run',role='settlement')
    assert len(d.tables[0].rows)==2
    assert d.tables[0].rows[0].source['kind']=='column_note'
    assert len(list(target_rows(d,d.tables[0],'结算时间')))==1
    assert not d.issues


def test_unapproved_mapping_rejected(tmp_path):
    mapping=tmp_path/'map.json';mapping.write_text(json.dumps({'approved':False,'mapping':{'x':'y'}}))
    with pytest.raises(DataError,match='approved'):TemplateRules(TEMPLATE,mapping)


def test_source_export_same_data_long_ids_and_formula_strings(tmp_path):
    source=tmp_path/'资金账单.csv'
    values=fund(备注='=HYPERLINK("https://example.com")')
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(FUND_COLUMNS);w.writerow([values.get(k) for k in FUND_COLUMNS])
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE,collection_context={'collection':[r for r in empty_collection('2026-08') if r['role']!='fund']})
    run=Path(result['run_dir']);report=Path(result['report'])
    assert report.name=='202608【抖音】Y-测试店.xlsx'
    assert json.loads((run/'manifest.json').read_text())['shop']['shop_name']=='测试店'
    with gzip.open(report.with_suffix('.xlsx.gz'),'rb') as f:assert f.read()==report.read_bytes()
    catalog=json.loads((run/'tables.json').read_text());raw=next(t for t in catalog['tables'] if t['role']=='fund')
    page=grid_page(run/'tables.sqlite',raw['table_id'],0,1)
    assert page['rows'][0]['values'][1]=="'1234567890123456789012345"
    assert len(page['columns'])==39
    wb=openpyxl.load_workbook(report,data_only=False,read_only=True)
    assert wb[raw['excel_sheet']]['AM2'].data_type=='s'
    assert wb[raw['excel_sheet']]['AM2'].value==values['备注']
    wb.close()
    cached=openpyxl.load_workbook(report,data_only=True,read_only=True)
    assert cached['核对汇总']['H2'].value==100.10
    assert cached['核对汇总']['H4'].value==100.10
    cached.close()
    manifest=json.loads((run/'manifest.json').read_text())
    assert manifest['tables']=='tables.json'
    assert all((run/a['path']).exists() for a in manifest['artifacts'])


def test_source_blank_steward_sheet_is_not_a_transaction(tmp_path):
    source=tmp_path/'empty.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='管家账户';ws['J7']='=4.02+0.96+14.5';wb.save(source)
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    schema=json.loads((Path(result['run_dir'])/'tables.json').read_text())
    assert not any(t['role']=='steward' for t in schema['tables'])


def test_grid_pagination_no_duplicate_or_loss(tmp_path):
    source=tmp_path/'fund.csv'
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(FUND_COLUMNS)
        for n in range(3):
            values=fund(动帐流水号=f'id{n}')
            w.writerow([values.get(k) for k in FUND_COLUMNS])
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    run=Path(result['run_dir']);catalog=json.loads((run/'tables.json').read_text());t=next(t for t in catalog['tables'] if t['role']=='fund')
    first=grid_page(run/'tables.sqlite',t['table_id'],0,2);second=grid_page(run/'tables.sqlite',t['table_id'],2,2)
    assert first['next_offset']==2 and second['next_offset'] is None
    assert len({r['row_id'] for r in first['rows']+second['rows']})==3


@pytest.mark.parametrize('approved',[False,True])
def test_excel_template_formulas_numeric_inputs_and_monthly_dependencies(tmp_path,approved):
    source=tmp_path/'fund.csv'
    rows=[fund(),fund(动账场景='提现',动账方向='出账',动账金额='5.00',订单实付应结=0,运费实付=0,平台服务费=0),
          fund(动账场景='测试未配置费用',动账金额='2.00',订单实付应结=0,运费实付=0,平台服务费=0)]
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(FUND_COLUMNS);w.writerows([[r.get(k) for k in FUND_COLUMNS] for r in rows])
    first=tmp_path/'settlement-a.csv';second=tmp_path/'settlement-b.csv'
    with first.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['订单号','结算金额','结算时间','子订单号'])
        w.writerows([['a','100.10','2026-08-01','a1'],['outside','888.00','2026-09-01','outside1']])
    with second.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(['结算时间','订单号','子订单号','结算金额'])
        w.writerows([['','','','收入+支出'],['2026-08-02','b','b1','3.10']])
    mapping=tmp_path/'approved.json';mapping.write_text(json.dumps({'approved':True,'mapping':{'测试未配置费用':'评价有礼'}}))
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'fund'},
                      {'path':str(first),'role':'settlement'},{'path':str(second),'role':'settlement'}]}))
    result=build([],manifest_path=manifest,month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE,mapping=mapping if approved else None)
    report=Path(result['report']);run=Path(result['run_dir'])
    wb=openpyxl.load_workbook(report,data_only=False,read_only=True)
    template=openpyxl.load_workbook(TEMPLATE,data_only=False,read_only=True)
    for row in template['核对汇总'].iter_rows():
        for c in row:
            if c.data_type=='f':
                actual=wb['核对汇总'][c.coordinate].value
                assert getattr(actual,'text',actual)==getattr(c.value,'text',c.value)
                if hasattr(c.value,'ref'):assert actual.ref==c.value.ref
    for coord in ['AN2','AO2','AP2']:
        original=getattr(template['资金账单 (核对)'][coord].value,'text',template['资金账单 (核对)'][coord].value)
        assert wb['资金账单 (核对)'][coord].value==original
    assert wb['资金账单 (核对)']['AN3'].value=='=_xlfn.IFS(C3="入账",D3,C3="出账",-D3,TRUE,D3)'
    assert wb['资金账单 (核对)']['AO4'].value=='=SUM(Q4:AB4)-R4'
    assert 'VLOOKUP(F4,账单字段对照表!A:B,2,FALSE)' in wb['资金账单 (核对)']['AP4'].value
    assert wb['资金账单 (核对)']['D2'].data_type=='n'
    assert wb['资金账单 (核对)']['Q2'].data_type=='n'
    assert wb['结算账单']['D1'].value=='结算金额'
    assert wb['结算账单'].max_row==3
    assert wb.sheetnames==template.sheetnames
    assert '结算账单原始' not in wb.sheetnames
    raw=next(t for t in json.loads((run/'tables.json').read_text())['tables'] if t['role']=='settlement')
    assert raw['excel_sheet'] is None
    assert grid_page(run/'tables.sqlite',raw['table_id'],0,2)['rows'][1]['values'][1]=='888.00'
    assert wb.calculation.fullCalcOnLoad
    template.close();wb.close()
    cached=openpyxl.load_workbook(report,data_only=True,read_only=True)
    assert cached['资金账单 (核对)']['AN3'].value==-5
    assert cached['资金账单 (核对)']['AP4'].value==('评价有礼' if approved else '#N/A')
    assert cached['核对汇总']['H3'].value==103.2
    assert {r[0]:r[1] for r in cached['账单字段对照表'].iter_rows(min_row=2,values_only=True)}.get('测试未配置费用')==('评价有礼' if approved else None)
    cached.close()
    page=grid_page(run/'tables.sqlite','processed_funds',0,3)
    assert page['rows'][0]['values'][39]=='100.10'
    assert page['rows'][2]['values'][41]==('评价有礼' if approved else None)
    from douyin_settle_recon.parse import import_file
    imported=Dataset({'month':'2026-08','shop_name':'测试店'})
    import_file(imported,report,tmp_path/'reimport')
    assert len(imported.table('settlement'))==1
    assert len(imported.table('settlement')[0].rows)==2
    process(imported,TemplateRules(TEMPLATE,mapping if approved else None))
    assert imported.checks[0]['实际值']==Decimal('103.20')


def test_supplied_sample_decimal_regression_and_screenshots(tmp_path):
    from douyin_settle_recon.parse import import_file
    original=Dataset({'month':'2026-08','shop_name':'样例','run_id':'regression'})
    import_file(original,ROOT/'doc/202608【抖音】Y-兹伊翰官方旗舰店-预期结果.xlsx',tmp_path)
    assert {t.role:len(t.rows) for t in original.tables if t.role in ['fund','settlement','deposit']}=={'fund':29896,'settlement':24201,'deposit':4}
    assert len(original.context['images'])==6
    references={r.values['坐标']:r.values['原始值'] for t in original.table('reference') for r in t.rows}
    assert references['F27']==28400.13 and references['J7']==19.48
    for image in original.context['images']:
        assert Path(image['path']).read_bytes().startswith(b'\x89PNG')
    for t in original.table('deposit'):
        for row in t.rows:row.source['account']='体验保证金账户'
    strict=Dataset(dict(original.context),tables=original.tables.copy())
    # Keep exercising the original 32-mapping fixture after the user updates
    # their editable template with these three formerly missing mappings.
    strict_rules=TemplateRules(TEMPLATE)
    for scene in ['评价有礼','以旧换新','快递拦截费']:
        strict_rules.mapping.pop(scene,None);strict_rules.snapshot['mapping'].pop(scene,None)
    processed=process(strict,strict_rules)
    assert len(processed.rows)==29900
    assert sum(i['code']=='unmapped_scene' for i in strict.issues)==48
    assert strict.checks[0]['实际值']==Decimal('1467515.09')
    assert strict.checks[1]['实际值']==Decimal('-7739.19')
    approved=tmp_path/'approved-test-mapping.json'
    approved.write_text(json.dumps({'approved':True,'mapping':{'评价有礼':'评价有礼','以旧换新':'货款收入-平台补贴','快递拦截费':'拦截费'}}))
    replay=Dataset(dict(original.context),tables=original.tables.copy())
    process(replay,TemplateRules(TEMPLATE,approved))
    assert replay.checks[0]['计算值']==replay.checks[0]['实际值']==Decimal('1467515.09')
    assert replay.checks[1]['计算值']==replay.checks[1]['实际值']==Decimal('-7739.19')
    assert not any(i['code']=='unmapped_scene' for i in replay.issues)
    assert replay.context['conclusion']=='待确认'  # historical balance evidence is still missing


def test_collection_failure_and_wrong_shop_are_never_reported_as_complete(tmp_path):
    source=tmp_path/'fund.csv'
    with source.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(FUND_COLUMNS);w.writerow([fund().get(k) for k in FUND_COLUMNS])
    manifest=tmp_path/'collection.json'
    payload={'shop':{'shopdesk_shop_uuid':'shop-a'},'inputs':[{'path':str(source),'role':'fund'}],
             'collection':[{'role':'settlement','status':'error','error':'permission denied'}]}
    manifest.write_text(json.dumps(payload))
    with pytest.raises(DataError,match='identity conflict'):
        build([],month='2026-08',shop_name='店',shopdesk_shop_uuid='shop-b',manifest_path=manifest,output=tmp_path/'wrong',template=TEMPLATE)
    result=build([],month='2026-08',shop_name='店',shopdesk_shop_uuid='shop-a',manifest_path=manifest,output=tmp_path/'right',template=TEMPLATE)
    data=json.loads((Path(result['run_dir'])/'result.json').read_text())
    assert result['conclusion']=='待确认'
    assert any(i['code']=='collection_failed' for i in data['issues'])
    missing=next(c for c in data['checks'] if c['项目']=='管家账户流水覆盖')
    assert missing['状态']=='待确认' and missing['实际值'] is None and missing['差异'] is None


def test_official_report_period_and_shop_binding():
    from douyin_settle_recon.collect import _period_matches,_report_ready
    good={'download_id':'DL202610181732251540952605','download_type':'CHARGE_BILL_ITEM',
          'start_time':'2026-09-01 00:00:00','end_time':'2026-09-30 23:59:59',
          'shop_id':'87755028','status':'SUCCESS'}
    assert _period_matches(good,'2026-09')
    assert not _period_matches(good,'2026-08')
    assert not _period_matches({**good,'end_time':'2026-09-29 23:59:59'},'2026-09')
    class FakeUI:
        def request(self,path):return {'code':0,'data':[good]}
        def wait(self,fn,timeout):return fn()
    manifest={'shop':{'platform_shop_id':'wrong-shop'}}
    from douyin_settle_recon.shopdesk import ShopDeskError
    with pytest.raises(ShopDeskError,match='different platform shop'):
        _report_ready(FakeUI(),good['download_id'],'2026-09',manifest)
    manifest={'shop':{}}
    assert _report_ready(FakeUI(),good['download_id'],'2026-09',manifest)==good
    assert manifest['shop']['platform_shop_id']=='87755028'


def test_official_wallet_signed_export_matches_retained_direction_formula(tmp_path):
    from douyin_settle_recon.parse import import_file
    source=tmp_path/'动账明细 - 管家账户.csv'
    rows=[['动账流水号','关联订单号','关联子订单号','动账时间','账户方向','动账金额(元)','冻结金额(元)','动账场景','账户余额(元)','冻结总额(元)','备注'],
          ['12345678901234567890','`12345','`12346','2026-09-01 12:00:00','支出','-6.00','0.00','小额打款','90.00','0.00','商家向消费者小额打款-补差价'],
          ['12345678901234567891','','','2026-09-02 12:00:00','收入','100.00','0.00','充值','190.00','0.00','管家账户充值'],
          ['12345678901234567892','','','2026-08-02 12:00:00','支出','-99.00','0.00','小额打款','90.00','0.00','消费者赔付']]
    with source.open('w',encoding='utf-8-sig',newline='') as f:csv.writer(f).writerows(rows)
    d=Dataset({'month':'2026-09','shop_name':'测试店','run_id':'wallet-test'})
    import_file(d,source,tmp_path/'run',role='steward',entity_name='测试主体')
    raw=d.table('steward')[0]
    p=process(d,TemplateRules(TEMPLATE))
    assert len(p.rows)==2
    debit,credit=[r.values for r in p.rows]
    assert debit['动帐流水号']=='12345678901234567890'
    assert debit['动账金额']==Decimal('6.00') and debit['动账方向']=='出账'
    assert debit['结算金额']==Decimal('-6.00') and debit['费用类型']=='消费者赔付'
    assert credit['动账金额']==Decimal('100.00') and credit['动账方向']=='入账'
    assert credit['结算金额']==Decimal('100.00')
    assert raw.rows[0].values['动账金额(元)']=='-6.00'
    assert raw.rows[0].values['账户方向']=='支出'
    assert any(i['code']=='outside_month' for i in d.issues)
