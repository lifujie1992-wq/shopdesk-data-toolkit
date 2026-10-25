from copy import deepcopy
from decimal import Decimal
from pathlib import Path
import json
import re
import zipfile
import xml.etree.ElementTree as ET

import openpyxl
import pytest

from douyin_settle_recon.model import DataError, sha256
from douyin_settle_recon.pipeline import build
from douyin_settle_recon.qianchuan import COLUMNS, statement
from douyin_settle_recon.report_validation import validate_report
from douyin_settle_recon.upload import prepare_metadata
from douyin_settle_recon.workbook_layout import NS, workbook_sheets
from recon_samples import empty_collection, q_receipt
from test_local_pipeline import TEMPLATE
from test_qianchuan import table


def sample(root,kind='empty',template=TEMPLATE):
    root.mkdir(parents=True,exist_ok=True)
    receipts=empty_collection();inputs=[];balances=[]
    if kind in ['data','detail_only','unproven_detail']:
        row=dict.fromkeys(COLUMNS,0)
        row.update({'日期':'2026-09-30','余额总消耗(元)':130.34,'非赠款消耗(元)':120.34,'赠款消耗(元)':10,'总存入(元)':1000,'总余额(元)':7,'赠款余额(元)':3,'非赠款余额(元)':4})
        total={'日期':'总计',**{k:row[k] for k in COLUMNS[1:11]},**dict.fromkeys(COLUMNS[11:],'-')}
        source=root/'source.xlsx';w=openpyxl.Workbook();s=w.active;s.append(COLUMNS)
        for r in ([total,row] if kind=='data' else [row]):s.append([r[k] for k in COLUMNS])
        w.save(source);w.close()
        inputs=[{'path':str(source.resolve()),'role':'qianchuan','entity_id':'123'}]
        receipts[-1]=q_receipt(1)
        if kind=='unproven_detail':receipts=receipts[:-1]
    elif kind=='failed':receipts[-1]={'role':'qianchuan','status':'permission_denied','error':'千川入口权限不足'}
    elif kind=='unknown_empty':receipts[-1]={'role':'qianchuan','status':'no_activity'}
    balances=[{'account':'千川账户','period':'2026-09','kind':'monthly','opening':'125','source':{'verified':True}}]
    manifest=root/'collection.json';manifest.write_text(json.dumps({'collection':receipts,'inputs':inputs,'balances':balances}))
    return build([],manifest_path=manifest,month='2026-09',shop_name='验收测试店',platform_shop_id='123',output=root/'runs',template=template)


def validation(run):
    root=Path(run['run_dir'])
    context=json.loads((root/'result.json').read_text())['context']
    snapshot=json.loads((root/'rules.json').read_text())
    return validate_report(run['report'],context,snapshot)


def mutate(path,sheet,change):
    path=Path(path)
    with zipfile.ZipFile(path) as z:
        parts={n:z.read(n) for n in z.namelist()};member=workbook_sheets(z)[sheet]
    xml=ET.fromstring(parts[member]);change(xml);parts[member]=ET.tostring(xml)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in parts.items():z.writestr(name,content)


def test_three_outcomes_and_valid_numeric_caches(tmp_path):
    for kind,status in [('data','success_data'),('empty','success_empty'),('failed','failed')]:
        run=sample(tmp_path/kind,kind);root=Path(run['run_dir']);result=validation(run)
        assert result['qianchuan_result']['status']==status
        assert result['ready']
        assert result['complete']==(kind!='failed')
        assert run['upload_ready']==result['ready']
        assert Path(run['report']).with_suffix('.xlsx.gz').exists()
        w=openpyxl.load_workbook(run['report'],data_only=True)
        if kind=='data':assert [w['核对汇总'][c].value for c in ['D44','D45','N6']]==[-120.34,990,869.66]
        if kind=='empty':
            assert [w['核对汇总'][c].value for c in ['D44','D45','N6']]==[0,0,0]
            assert w['核对汇总']['M6'].value==125 and w['核对汇总']['O6'].value==125
            assert w['千川账户']['Q4'].value=='#N/A' and result['balance_issues']
        w.close()
        if kind=='failed':
            assert result['expense_status']=='partial' and run['partial_accounts']==['千川账户']
        metadata,report,gz=prepare_metadata(root)
        assert metadata['report']['sha256']==sha256(report) and metadata['upload']['sha256']==sha256(gz)
        assert not {'qianchuan_result','expense_status','upload_ready','partial_accounts'} & set(metadata)
        if kind=='failed':assert metadata['reconciliation_status']=='pending'


def test_cli_failed_collection_is_not_reported_as_completed(tmp_path,capsys):
    from douyin_settle_recon.cli import main
    receipts=empty_collection()
    receipts[-1]={'role':'qianchuan','status':'error','error':'千川查询超时'}
    source=tmp_path/'collection.json'
    source.write_text(json.dumps({'collection':receipts}))
    assert main(['import','--manifest',str(source),'--month','2026-09',
                 '--shop-name','验收测试店','--output',str(tmp_path/'runs'),
                 '--template',str(TEMPLATE)])==3
    result=json.loads(capsys.readouterr().out)
    assert result['upload_ready'] is False and Path(result['validation']).is_file()


@pytest.mark.parametrize('coord',['D2','H2','D44','D45','N6','O6','P6'])
def test_partial_delivery_keeps_other_checks_and_rejects_fabricated_qianchuan_zero(tmp_path,coord):
    run=sample(tmp_path,'failed')
    def change(xml):
        c=xml.find(f".//m:c[@r='{coord}']",NS);v=c.find('m:v',NS)
        if v is None:v=ET.SubElement(c,'{'+NS['m']+'}v')
        if coord in ['D2','H2']:v.text=None
        else:c.set('t','n');v.text='0'
    mutate(run['report'],'核对汇总',change)
    result=validation(run)
    assert not result['ready'] and any(coord in e['cell'] for e in result['errors'])
    with pytest.raises(DataError,match='费用校验失败'):prepare_metadata(Path(run['run_dir']))


def test_permission_receipt_alone_cannot_skip_failed_other_sources(tmp_path):
    run=sample(tmp_path,'failed');root=Path(run['run_dir'])
    context=json.loads((root/'result.json').read_text())['context']
    context['collection'].append({'role':'fund','status':'error','error':'下载失败'})
    assert not validate_report(run['report'],context)['ready']
    context['collection'].pop();context.pop('qianchuan_access')
    assert not validate_report(run['report'],context)['ready']


def test_partial_report_is_pending_even_if_manifest_claims_matched(tmp_path):
    run=sample(tmp_path,'failed');root=Path(run['run_dir'])
    path=root/'manifest.json';manifest=json.loads(path.read_text());manifest['conclusion']='完全一致'
    path.write_text(json.dumps(manifest))
    metadata,_,_=prepare_metadata(root)
    assert metadata['reconciliation_status']=='pending'


def test_only_verified_complete_details_can_generate_total(tmp_path):
    run=sample(tmp_path/'complete','detail_only');assert validation(run)['ready']
    w=openpyxl.load_workbook(run['report'],data_only=True)
    assert sum(row[0]=='总计' for row in w['千川账户'].values)==1
    assert w['核对汇总']['D45'].value==990;w.close()
    assert not validation(sample(tmp_path/'unproven','unproven_detail'))['ready']
    assert not validation(sample(tmp_path/'unknown','unknown_empty'))['ready']


def test_missing_qianchuan_sheet_cannot_bypass_source_validation(tmp_path):
    run=sample(tmp_path)
    path=Path(run['report'])
    with zipfile.ZipFile(path) as z:parts={n:z.read(n) for n in z.namelist()}
    xml=ET.fromstring(parts['xl/workbook.xml']);sheets=xml.find('m:sheets',NS)
    sheets.remove(next(s for s in sheets if s.get('name')=='千川账户'))
    parts['xl/workbook.xml']=ET.tostring(xml)
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as z:
        for name,content in parts.items():z.writestr(name,content)
    result=validation(run)
    assert not result['ready'] and any(e['cell']=='千川账户' for e in result['errors'])


@pytest.mark.parametrize('coord,kind',[('D44','empty'),('D45','empty'),('D45','str'),('D10','bool'),('D10','nan'),('D10','inf'),('D10','error'),('N6','wrong'),('H2','wrong'),('H3','wrong'),('H4','empty')])
def test_invalid_final_caches_block_upload(tmp_path,coord,kind):
    run=sample(tmp_path)
    def change(xml):
        c=xml.find(f".//m:c[@r='{coord}']",NS);v=c.find('m:v',NS)
        if kind=='empty':v.text=None
        elif kind=='str':c.set('t','str');v.text='0'
        elif kind=='bool':c.set('t','b');v.text='0'
        elif kind=='error':c.set('t','e');v.text='#N/A'
        else:v.text={'nan':'NaN','inf':'Infinity','wrong':'99'}[kind]
    mutate(run['report'],'核对汇总',change)
    report=validation(run)
    assert not report['ready'] and any(coord in e['cell'] for e in report['errors'])
    with pytest.raises(DataError,match='费用校验失败'):prepare_metadata(Path(run['run_dir']))


@pytest.mark.parametrize('kind',['no_total','two_totals','wrong_header','missing_amount','half_name','formula_changed','array_lost'])
def test_source_and_formula_integrity_in_final_file(tmp_path,kind):
    run=sample(tmp_path)
    def change(xml):
        if kind=='two_totals':
            row=deepcopy(xml.find(".//m:row[@r='2']",NS));row.set('r','10')
            for c in row:c.set('r',re.sub(r'\d+$','10',c.get('r')))
            xml.find('m:sheetData',NS).append(row)
        elif kind in ['no_total','wrong_header','missing_amount']:
            coord={'no_total':'A2','wrong_header':'C1','missing_amount':'C2'}[kind]
            c=xml.find(f".//m:c[@r='{coord}']",NS)
            for node in list(c):c.remove(node)
        elif kind=='half_name':
            c=xml.find(".//m:c[@r='C10']",NS)
            for node in list(c):c.remove(node)
        else:
            f=xml.find(".//m:c[@r='D45']/m:f",NS)
            if kind=='array_lost':f.attrib.clear()
            else:f.text='IFERROR(1,0)'
    mutate(run['report'],'千川账户' if kind in ['no_total','two_totals','wrong_header','missing_amount'] else '核对汇总',change)
    assert not validation(run)['ready']


def test_qianchuan_fees_are_located_by_names_when_rows_move(tmp_path):
    path=tmp_path/'moved.xlsx';w=openpyxl.load_workbook(TEMPLATE);s=w['核对汇总']
    for old,new in [(44,48),(45,49)]:
        for col in ['B','C','D']:
            v=deepcopy(s[f'{col}{old}'].value)
            if hasattr(v,'ref'):v.ref=f'{col}{new}'
            s[f'{col}{new}']=v;s[f'{col}{old}']=None
    w.save(path);w.close()
    run=sample(tmp_path/'run',template=path)
    assert validation(run)['ready']
    w=openpyxl.load_workbook(run['report'],data_only=True)
    assert w['核对汇总']['D48'].value==w['核对汇总']['D49'].value==0
    w.close()


@pytest.mark.parametrize('zero_total',[False,True])
def test_header_only_file_needs_matching_successful_empty_query(tmp_path,zero_total):
    root=tmp_path/'case';root.mkdir();source=root/'empty.xlsx'
    w=openpyxl.Workbook();w.active.append(COLUMNS)
    if zero_total:w.active.append(['总计',*([0]*10),'-','-','-'])
    w.save(source);w.close()
    manifest=root/'collection.json'
    manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'qianchuan','entity_id':'123'}],
                                   'collection':empty_collection()}))
    run=build([],manifest_path=manifest,month='2026-09',shop_name='空查询',output=root/'runs',template=TEMPLATE)
    assert validation(run)['ready']
    receipts=empty_collection();receipts[-1]['query']['params']['start_day']=['2026-08-01']
    manifest.write_text(json.dumps({'collection':receipts}))
    run=build([],manifest_path=manifest,month='2026-09',shop_name='错误月份',output=root/'runs',template=TEMPLATE)
    assert not validation(run)['ready']


def test_total_beyond_10000_rows_is_checked_without_scan_cutoff(tmp_path):
    run=sample(tmp_path)
    def change(xml):
        row=xml.find(".//m:row[@r='2']",NS);data=xml.find('m:sheetData',NS);data.remove(row)
        row.set('r','10005')
        for c in row:c.set('r',re.sub(r'\d+$','10005',c.get('r')))
        data.append(row)
    mutate(run['report'],'千川账户',change)
    assert validation(run)['ready']
