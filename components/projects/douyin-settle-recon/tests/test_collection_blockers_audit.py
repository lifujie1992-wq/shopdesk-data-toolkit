import importlib.util
import json
from pathlib import Path
import sys

from openpyxl import load_workbook
from PIL import Image
import pytest

from douyin_settle_recon.model import DataError

TOOLS = Path(__file__).resolve().parents[1] / 'tools'
sys.path.insert(0, str(TOOLS))
import audit_collection_blockers as audit
import render_collection_blockers as render


@pytest.mark.parametrize('message', ['暂无权限', '无访问权限', '权限不足', '所属的岗位没有访问该页面的权限',
                                     '抱歉，你的权限被主账号调整无法访问该页面\n权限点:「流量-千川推广」'])
def test_explicit_official_denial_is_reported_instead_of_generic_timeout(message):
    code, reason = audit.classify('finance UI condition did not become ready', {'text': message, 'url': audit.ENTRY})
    assert code == 'permission' and all(line in reason for line in message.splitlines())


def test_no_activity_and_unopened_accounts_are_not_classified_as_permission_denials():
    assert audit.restriction({'text': '暂无数据\n开通抖店钱包\n去开通', 'url': 'https://fxg.jinritemai.com/ffa/fund-control/butler-account'}) is None
    assert audit.classify(DataError('wrong downloaded schema'), {'text': '日结报表'})[0] == 'validation'
    assert audit.classify('finance UI condition did not become ready', {'text': '加载中'})[0] == 'unconfirmed'


def fixture_audit(tmp_path):
    root = tmp_path / 'audit'
    directory = root / 'shops' / 'uuid-1'
    directory.mkdir(parents=True)
    shop = {'id': 'uuid-1', 'name': '=店铺名称'}
    screenshot = directory / 'blocked.png'
    Image.new('RGB', (400, 240), 'white').save(screenshot)
    roles = [{'role': role, 'complete': role != 'qianchuan',
              'receipts': [{'role': role, 'status': 'error' if role == 'qianchuan' else 'not_opened' if role == 'steward' else 'downloaded'}],
              'downloaded_files': []} for role in audit.PAGES]
    issue = {'role': 'qianchuan', 'status': 'permission', 'reason': '暂无权限', 'checked_at': '2026-10-22', 'screenshot': str(screenshot)}
    report = {'shop': shop, 'complete': True, 'started_at': '2026-10-22', 'finished_at': '2026-10-22', 'roles': roles, 'issues': [issue]}
    (root / 'shops.json').write_text(json.dumps({'month': '2026-09', 'shops': [shop]}))
    (directory / 'audit.json').write_text(json.dumps(report))
    (directory / 'collection.json').write_text(json.dumps({'screenshots': []}))
    return root, directory, report


def test_incomplete_real_shop_scope_cannot_be_rendered_as_complete(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    report['roles'].pop()
    (directory / 'audit.json').write_text(json.dumps(report))
    with pytest.raises(ValueError, match='incomplete'):
        render.render(root, tmp_path / 'result.xlsx')
    assert not (tmp_path / 'result.xlsx').exists()


def test_customer_report_embeds_evidence_and_preserves_plain_text_shop_names(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    output = tmp_path / 'result.xlsx'
    result = render.render(root, output)
    assert result['shops'] == 1 and result['source_checks'] == 7
    assert result['blocked_shops'] == 1 and result['issue_count'] == 1
    assert result['unopened_accounts'] == 1 and result['embedded_screenshots'] == 1
    wb = load_workbook(output)
    cell = wb['抓取阻塞清单']['A2']
    assert cell.value == '=店铺名称' and cell.data_type == 's'
    assert '千川推广' in wb['抓取阻塞清单']['E2'].value
    assert wb['抓取阻塞清单']['H2'].hyperlink.location.startswith("'截图证据'!")
    assert len(wb['截图证据']._images) == 1
    wb.close()


def test_transient_retries_preserve_explicit_permission_denials_and_successful_subjects(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    report['issues'].append({'role': 'tax', 'status': 'unconfirmed', 'reason': 'query confirmation timeout'})
    report['roles'][5]['complete'] = False
    report['roles'][5]['receipts'] = [{'role': 'tax', 'status': 'downloaded', 'entity': {'name': '主体一'}}, {'role': 'tax', 'status': 'error', 'entity': {'name': '主体二'}}]
    (directory / 'audit.json').write_text(json.dumps(report))
    (directory / 'collection.json').write_text(json.dumps({'collection': report['roles'][5]['receipts'] + report['roles'][-1]['receipts']}))
    audit.prepare_transient_retries(root, [report['shop']])
    updated = json.loads((directory / 'audit.json').read_text())
    manifest = json.loads((directory / 'collection.json').read_text())
    assert updated['complete'] is False
    assert all(x['role'] != 'tax' for x in updated['roles'])
    assert updated['issues'] == [report['issues'][0]]
    assert any(x['role'] == 'qianchuan' for x in updated['roles'])
    assert [x['status'] for x in manifest['collection'] if x['role'] == 'tax'] == ['downloaded']
    assert (directory / 'attempt-history/1.json').is_file()


def test_incomplete_dates_are_explained_from_the_download_without_inferring_missing_amounts(tmp_path):
    from openpyxl import Workbook
    root, directory, report = fixture_audit(tmp_path)
    source = directory / 'daily.xlsx'
    wb = Workbook()
    wb.active.append(['日期'])
    wb.active.append(['2026-09-30'])
    wb.active.append(['2026-09-20'])
    wb.save(source)
    issue = report['issues'][0]
    issue.update(status='unconfirmed', advertiser_id='123', collector_error='千川账户日结报表日期不完整')
    manifest = {'inputs': [{'role': 'qianchuan', 'advertiser_id': '123', 'path': str(source)}]}
    (directory / 'collection.json').write_text(json.dumps(manifest))
    audit.explain_incomplete_dates(root, [report], '2026-09')
    assert issue['status'] == 'validation'
    assert len(issue['date_validation']['missing_dates']) == 28
    assert issue['date_validation']['actual_dates'] == ['2026-09-20', '2026-09-30']
    assert '程序' in issue['reason'] and '金额' in issue['action']


def test_closed_shop_without_target_month_is_explained_without_claiming_zero_balance(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    issue = report['issues'][0]
    issue.update(role='fund', status='unconfirmed', page_excerpt='财气满满服饰店\n已清退\n月汇总\n2026-06\t73\t资金金额\n2026-05\t10')
    (directory / 'collection.json').write_text(json.dumps({'inputs': []}))
    audit.explain_incomplete_dates(root, [report], '2026-09')
    assert issue['status'] == 'period_missing'
    assert issue['visible_statement_months'] == ['2026-05', '2026-06']
    assert '不能' in issue['reason'] and '0' in issue['reason']


def test_pixel_identical_permission_screens_keep_separate_shop_evidence_captions(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    other = json.loads(json.dumps(report))
    other['shop'] = {'id': 'uuid-2', 'name': '另一家店铺'}
    other_dir = root / 'shops/uuid-2'
    other_dir.mkdir()
    (other_dir / 'audit.json').write_text(json.dumps(other))
    (other_dir / 'collection.json').write_text(json.dumps({'screenshots': []}))
    (root / 'shops.json').write_text(json.dumps({'month': '2026-09', 'shops': [report['shop'], other['shop']]}))
    output = tmp_path / 'two-shops.xlsx'
    result = render.render(root, output)
    assert result['embedded_screenshots'] == 2
    wb = load_workbook(output)
    links = [wb['抓取阻塞清单'][f'H{r}'].hyperlink.location for r in (2, 3)]
    assert links[0] != links[1]
    for location, expected in zip(links, [report['shop']['name'], other['shop']['name']]):
        assert wb['截图证据'][location.split('!')[1]].value.startswith(expected)
    wb.close()


@pytest.mark.parametrize(('total', 'expected'), [(2, 'period_missing'), (25, 'unconfirmed')])
def test_missing_month_on_an_open_shop_requires_a_complete_visible_month_list(tmp_path, total, expected):
    root, directory, report = fixture_audit(tmp_path)
    issue = report['issues'][0]
    issue.update(role='fund', status='unconfirmed', page_excerpt=f'正常营业\n月汇总\n2026-06\t10\n2026-05\t20\n共 {total} 条')
    (directory / 'collection.json').write_text(json.dumps({'inputs': []}))
    audit.explain_incomplete_dates(root, [report], '2026-09')
    assert issue['status'] == expected


def test_qianchuan_registration_form_is_not_reported_as_a_permission_denial(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    issue = report['issues'][0]
    issue.update(status='unconfirmed', actual_url='https://fxg.jinritemai.com/ffa/ad/promotion-v2/account/register',
                 page_excerpt='免费开通巨量千川账户\n账户名称\n主体名称\n立即免费开户')
    (directory / 'collection.json').write_text(json.dumps({'inputs': []}))
    audit.explain_incomplete_dates(root, [report], '2026-09')
    assert issue['status'] == 'qianchuan_not_opened'
    assert '开户' in issue['required_page_path']


def test_empty_deposit_failure_requires_an_observed_placeholder_and_zero_count(tmp_path):
    root, directory, report = fixture_audit(tmp_path)
    issue = report['issues'][0]
    issue.update(role='base_deposit', status='download')
    (directory / 'collection.json').write_text(json.dumps({'inputs': []}))
    (directory / 'empty-deposit-proof.json').write_text(json.dumps({'rows': [['暂无数据']], 'text': '共 0 条\n暂无数据',
                                                                  'screenshot': issue['screenshot']}))
    audit.explain_incomplete_dates(root, [report], '2026-09')
    assert issue['status'] == 'empty_table'
    assert '占位行' in issue['reason'] and '开发' in issue['action']
