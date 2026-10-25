from copy import copy

import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from PIL import Image as PILImage
import pytest

from douyin_settle_recon.pipeline import build
from test_local_pipeline import TEMPLATE, FUND_COLUMNS, fund


def test_report_keeps_template_names_source_styles_anchors_and_reference_formulas(tmp_path):
    image_path=tmp_path/'capture.png'
    PILImage.new('RGB',(180,90),'#336699').save(image_path)
    source=tmp_path/'source.xlsx'
    wb=openpyxl.Workbook();ws=wb.active;ws.title='资金账单'
    ws.append(FUND_COLUMNS);row=fund();ws.append([row.get(k) for k in FUND_COLUMNS])
    ws.append([row.get(k) for k in FUND_COLUMNS]);ws['D3'].fill=PatternFill('solid',fgColor='33CC66')
    ws['A1'].font=Font(name='Arial',size=13,color='123456',bold=True)
    ws['A1'].fill=PatternFill('solid',fgColor='EEDDCC')
    ws['A1'].alignment=Alignment(horizontal='center',wrap_text=True)
    ws['A1'].border=Border(bottom=Side(style='thin',color='224466'))
    ws.column_dimensions['A'].width=23.4;ws.row_dimensions[1].height=31;ws.freeze_panes='C2'
    ws.page_setup.orientation='landscape';ws.page_setup.paperSize=ws.PAPERSIZE_A3
    ws.oddHeader.center.text='月度账单';ws.sheet_properties.tabColor='336699'
    ws=wb.create_sheet('余额账单');ws['F27']=28400.13;ws.column_dimensions['F'].width=18
    ws=wb.create_sheet('管家账户');ws['J7']='=4.02+0.96+14.5'
    deposit=wb.create_sheet('保证金账户')
    deposit.append(['操作单号','操作时间','操作类型','操作金额','备注','关联订单号','交易状态'])
    deposit.append(['deposit-01','2026-08-01','充值',10.5,'保证金充值','order-01','成功'])
    for name,count in [('余额账单',1),('保证金账户',2),('管家账户',3)]:
        for i in range(count):
            image=Image(image_path)
            image.anchor=TwoCellAnchor(editAs='oneCell',_from=AnchorMarker(col=8 if name=='保证金账户' else 0,row=i*18,colOff=19685,rowOff=635),
                                      to=AnchorMarker(col=16 if name=='保证金账户' else 8,row=i*18+15,colOff=33020,rowOff=137160))
            wb[name].add_image(image)
    summary=wb.create_sheet('核对汇总');summary['A2']='example-code';summary['M2']=30000
    summary['P2']='=+O2-余额账单!F27';summary['P3']='=+O3-管家账户!J7'
    wb.save(source);wb.close()
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    import zipfile
    with zipfile.ZipFile(result['report']) as generated,zipfile.ZipFile(TEMPLATE) as original:
        assert generated.read('xl/theme/theme1.xml')==original.read('xl/theme/theme1.xml')
    template=openpyxl.load_workbook(TEMPLATE,read_only=True)
    actual=openpyxl.load_workbook(result['report'])
    expected=openpyxl.load_workbook(source)
    assert actual.sheetnames==template.sheetnames
    assert actual.active.title=='核对汇总'
    for attribute in ['font','fill','border','alignment','number_format']:
        assert copy(getattr(actual['资金账单']['A1'],attribute))==copy(getattr(expected['资金账单']['A1'],attribute))
    assert copy(actual['资金账单']['D3'].fill)==copy(expected['资金账单']['D3'].fill)
    assert actual['资金账单'].column_dimensions['A'].width==23.4
    assert actual['资金账单'].row_dimensions[1].height==31
    assert actual['资金账单'].freeze_panes=='C2'
    assert actual['资金账单'].page_setup.orientation=='landscape'
    assert actual['资金账单'].page_setup.paperSize==expected['资金账单'].page_setup.paperSize
    assert actual['资金账单'].oddHeader.center.text=='月度账单'
    assert actual['资金账单'].sheet_properties.tabColor==expected['资金账单'].sheet_properties.tabColor
    assert actual['管家账户']['J7'].value=='=4.02+0.96+14.5'
    assert actual['余额截图']['F27'].value==28400.13
    assert actual['核对汇总']['A2'].value=='example-code'
    assert actual['核对汇总']['M2'].value==30000
    assert actual['核对汇总']['P2'].value=='=+O2-余额截图!F27'
    assert actual['核对汇总']['P3'].value=='=+O3-管家账户!J7'
    assert actual['核对汇总']['M2'].comment
    assert [actual['资金账单 (核对)'][coord].value for coord in ['A4','B4','I4']]==[None,None,None]
    assert actual['资金账单 (核对)']['D4'].value==10.5
    assert actual['保证金账户']['A2'].value=='deposit-01'
    from douyin_settle_recon.export import grid_page
    from pathlib import Path
    row=grid_page(Path(result['run_dir'])/'tables.sqlite','processed_funds',2,1)['rows'][0]['values']
    assert row[0]=='2026-08-01' and row[1]=='deposit-01' and row[8]=='order-01'
    for name,source_name in [('余额截图','余额账单'),('保证金账户','保证金账户'),('管家账户','管家账户')]:
        assert len(actual[name]._images)==len(expected[source_name]._images)
        for got,want in zip(actual[name]._images,expected[source_name]._images):
            assert got.anchor._from==want.anchor._from
            assert got.anchor.to==want.anchor.to
            assert got.anchor.editAs==want.anchor.editAs
            assert got._data()==want._data()
    for w in [actual,expected,template]:w.close()


def test_live_captures_use_account_sheets_and_do_not_overlap_tables(tmp_path):
    import csv,json
    from pathlib import Path
    source=tmp_path/'fund.csv'
    with source.open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(FUND_COLUMNS);writer.writerow([fund().get(k) for k in FUND_COLUMNS])
    labels=['资金账单月汇总','结算账单月汇总','基础保证金账户余额','体验保证金账户余额','管家账户余额','涉税账单-主体']
    shots=[]
    for i,label in enumerate(labels):
        path=tmp_path/f'capture-{i}.png';PILImage.new('RGB',(600,300+i*60)).save(path)
        shots.append({'path':str(path),'sheet':label})
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'inputs':[{'path':str(source)}],'screenshots':shots}))
    result=build([],manifest_path=manifest,month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    actual=openpyxl.load_workbook(result['report'])
    assert {s.title:len(s._images) for s in actual if s._images}=={'结算账单':1,'保证金账户':2,'管家账户':1,'余额截图':1}
    assert not actual['资金账单']._images
    assert actual['余额截图']._images[0].anchor._from.col==0
    handoff=json.loads((Path(result['run_dir'])/'manifest.json').read_text())
    assert next(s for s in handoff['screenshots'] if '涉税' in s['sheet'])['excel_sheet'] is None
    assert actual['保证金账户']._images[1].anchor._from.row>actual['保证金账户']._images[0].anchor.to.row
    actual.close()


def test_settlement_source_columns_keep_spelling_and_order(tmp_path):
    import csv
    source=tmp_path/'settlement.csv'
    headers=['结算时间','订单号','子订单号','结算金额','结算账户','店铺券','APP渠道','备注']
    with source.open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(headers);writer.writerow(['2026-08-01','01','02','12.30','货款账户','20','其他','备注原文'])
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    actual=openpyxl.load_workbook(result['report'],read_only=True,data_only=True)
    ws=actual['结算账单']
    names=[c.value for c in next(ws.iter_rows(max_row=1))]
    assert names[:len(headers)]==headers
    assert '店铺劵' not in names
    assert ws['F2'].value==20
    assert ws['G2'].value=='其他'
    assert ws['H2'].value=='备注原文'
    assert actual['核对汇总']['H3'].value==12.3
    actual.close()


def test_verified_monthly_balance_is_written_next_to_its_screenshot_reference(tmp_path):
    import csv
    source=tmp_path/'fund.csv'
    with source.open('w',newline='') as stream:
        writer=csv.writer(stream);writer.writerow(FUND_COLUMNS);writer.writerow([fund().get(k) for k in FUND_COLUMNS])
    balance={'account':'货款账户','period':'2026-08','kind':'monthly','opening':'10','closing':'110.10','source':{'official_month_row':'2026-08'}}
    result=build([source],month='2026-08',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE,balances=[balance])
    book=openpyxl.load_workbook(result['report'],read_only=True,data_only=False)
    assert book['余额截图']['F27'].value==110.1
    assert book['核对汇总']['P2'].value=='=+O2-余额截图!F27'
    book.close()
    cached=openpyxl.load_workbook(result['report'],read_only=True,data_only=True)
    assert cached['核对汇总']['P2'].value==0
    cached.close()


def test_wallet_screenshots_have_subject_list_and_editable_total_without_double_counting(tmp_path):
    import json
    from pathlib import Path
    image=tmp_path/'wallet.png';PILImage.new('RGB',(1920,1080)).save(image)
    shots=[]
    for i,amount in enumerate(['4.02','0.96','14.50','0.00',None]):
        shots.append({'path':str(image),'sheet':'管家账户余额','entity_name':f'主体{i+1}',
                      'captured_at':'2026-10-19T08:36:00Z',
                      'wallet_balance':{'status':'not_opened' if amount is None else 'opened','total':amount,'kind':'current'}})
    # The earlier screenshot remains evidence but cannot add a second balance.
    shots.insert(0,{**shots[0],'captured_at':'2026-10-19T08:35:00Z',
                    'wallet_balance':{'status':'opened','total':'100.00','kind':'current'}})
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'screenshots':shots}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report']);ws=book['管家账户']
    assert ws['J7'].value=='=SUM(J11:J15)'
    assert [ws[f'J{r}'].value for r in range(11,16)]==[4.02,.96,14.5,0,None]
    assert ws['K15'].value=='未开通'
    assert ws['L11'].value=='2026-10-19 16:36:00'
    assert '非月末余额' in ws['I3'].value
    assert len(ws._images)==6
    assert all(im.anchor.to.col<8 for im in ws._images)
    # Current screenshots never become verified September closing balances.
    assert book['核对汇总']['P3'].value is None
    book.close()
    cached=openpyxl.load_workbook(result['report'],read_only=True,data_only=True)
    assert cached['管家账户']['J7'].value==19.48
    cached.close()
    handoff=json.loads((Path(result['run_dir'])/'manifest.json').read_text())
    assert handoff['screenshots'][1]['wallet_balance']['total']=='4.02'


def test_wallet_missing_balance_cannot_be_presented_as_complete_total():
    from douyin_settle_recon.export import _wallet_balance_grid
    from douyin_settle_recon.model import Dataset
    ds=Dataset({'images':[
        {'sheet':'管家账户','entity_name':'已采集','wallet_balance':{'total':'4.02','kind':'current'}},
        {'sheet':'管家账户','entity_name':'未采集'},
        {'sheet':'管家账户','entity_name':'已采集','captured_at':'2026-10-19T08:36:00Z'},
    ]})
    grid=_wallet_balance_grid(ds)
    assert grid[6][9][1]=='#N/A'
    assert 'COUNT(J11:J12)=2' in grid[6][9][0]
    assert grid[10][9] is None  # Do not reuse an older balance when latest capture lacks it.


@pytest.mark.parametrize('extension',['csv','xlsx'])
def test_wallet_ledger_is_used_for_calculation_but_never_displayed_in_wallet_sheet(tmp_path,extension):
    import csv,json
    from pathlib import Path
    headers=['动账流水号','关联订单号','关联子订单号','动账时间','账户方向','动账金额(元)',
             '冻结金额(元)','动账场景','账户余额(元)','冻结总额(元)','备注']
    rows=[[f'1234567890123456789{i}','','','2026-09-02 12:00:00','收入','10.00','0.00',
           '充值','190.00','0.00','管家账户充值'] for i in range(8)]
    source=tmp_path/f'管家账户.{extension}'
    if extension=='csv':
        with source.open('w',encoding='utf-8-sig',newline='') as stream:
            writer=csv.writer(stream);writer.writerow(headers);writer.writerows(rows)
    else:
        wb=openpyxl.Workbook();ws=wb.active;ws.title='管家账户';ws.append(headers)
        for row in rows:ws.append(row)
        wb.save(source);wb.close()
    image=tmp_path/'wallet.png';PILImage.new('RGB',(1920,1080),'#336699').save(image)
    shots=[{'path':str(image),'sheet':'管家账户余额','entity_name':'主体一','captured_at':'2026-10-23T10:00:00Z',
            'wallet_balance':{'status':'opened','total':'19.48','kind':'current'}}]
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'steward','entity_name':'主体一'}],
                                                                'screenshots':shots}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report']);ws=book['管家账户']
    values=[c.value for row in ws for c in row if c.value is not None]
    assert not set(headers)&set(values)
    assert not any(row[0] in values for row in rows)
    assert ws['I1'].value=='管家账户余额汇总' and ws['I7'].value=='总余额（元）'
    assert [ws.cell(10,c).value for c in range(9,13)]==['账户主体','总余额（元）','账户状态','截图时间（北京时间）']
    assert ws['I11'].value=='主体一' and ws['J11'].value==19.48 and ws['K11'].value=='已开通'
    assert ws['L11'].value=='2026-10-23 18:00:00' and ws['J7'].value=='=SUM(J11:J11)'
    assert len(ws._images)==1 and ws._images[0]._data()==image.read_bytes()
    assert ws._images[0].anchor._from.col==0 and ws._images[0].anchor.to.col<8
    assert all(c.value is None for row in ws.iter_rows(max_col=8) for c in row)
    book.close()
    cached=openpyxl.load_workbook(result['report'],data_only=True)
    assert cached['管家账户']['J7'].value==19.48
    cached.close()
    data=json.loads((Path(result['run_dir'])/'tables.json').read_text())
    raw=next(t for t in data['tables'] if t['role']=='steward')
    assert raw['excel_sheet'] is None
    from douyin_settle_recon.export import grid_page
    saved=grid_page(Path(result['run_dir'])/'tables.sqlite',raw['table_id'],0,10)
    assert len(saved['rows'])==8 and saved['rows'][0]['values'][0]==rows[0][0]
    outcome=json.loads((Path(result['run_dir'])/'result.json').read_text())
    coverage=next(c for c in outcome['checks'] if c['项目']=='管家账户流水覆盖')
    assert coverage['实际值']=='80.00' and coverage['计算值']=='80.00' and coverage['差异']=='0.00'


def test_imported_wallet_summary_keeps_subjects_used_by_its_total_formula(tmp_path):
    source=tmp_path/'summary.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.title='管家账户'
    ws['A1']='动账流水号';ws['A2']='不应显示的旧流水'
    ws['I1']='管家账户余额汇总';ws['I7']='总余额（元）';ws['J7']='=SUM(J11:J11)'
    for col,value in zip(['I','J','K','L'],['账户主体','总余额（元）','账户状态','截图时间（北京时间）']):ws[f'{col}10']=value
    for col,value in zip(['I','J','K','L'],['旧报表主体',4.02,'已开通','2026-10-23 18:00:00']):ws[f'{col}11']=value
    wb.save(source);wb.close()
    result=build([source],month='2026-09',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report']);ws=book['管家账户']
    assert ws['A1'].value is None and ws['A2'].value is None
    assert ws['J7'].value=='=SUM(J11:J11)' and ws['J11'].value==4.02 and ws['I11'].value=='旧报表主体'
    book.close()

