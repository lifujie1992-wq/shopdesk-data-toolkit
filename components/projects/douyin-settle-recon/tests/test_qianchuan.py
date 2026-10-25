from recon_samples import q_receipt
from copy import deepcopy
from datetime import date, timedelta
from decimal import Decimal
import json
from pathlib import Path

import openpyxl
from PIL import Image
import pytest

from douyin_settle_recon.model import DataError, Table
from douyin_settle_recon.pipeline import build
from douyin_settle_recon.qianchuan import COLUMNS, TOTAL_FORMULAS, statement
from test_local_pipeline import TEMPLATE


def records(month='2026-09',days=30):
    daily=[]
    for i in range(days):
        row=dict.fromkeys(COLUMNS,Decimal(0))
        row.update({'日期':date.fromisoformat(month+'-01')+timedelta(days=i),
                    '余额总消耗(元)':Decimal('2.10'),'非赠款消耗(元)':Decimal('2'),
                    '赠款消耗(元)':Decimal('.10'),'总存入(元)':Decimal('3'),
                    '总余额(元)':Decimal(i+1),'赠款余额(元)':Decimal('.20'),
                    '非赠款余额(元)':Decimal(i+1)-Decimal('.20')})
        daily.append(row)
    total={'日期':'总计',**{k:sum(r[k] for r in daily) for k in COLUMNS[1:11]},**{k:'-' for k in COLUMNS[11:]}}
    return [total,*daily]


def table(rows):return Table.from_records('daily','日结报表',rows,role='qianchuan',fields=COLUMNS)


@pytest.mark.parametrize('reverse',[False,True])
def test_month_end_uses_calendar_date_and_decimal_totals(reverse):
    rows=records()
    if reverse:rows=[rows[0],*reversed(rows[1:])]
    view,values=statement([table(rows)],'2026-09')
    assert values['consumption']==Decimal('-60') and values['recharge']==Decimal('87')
    assert values['grant']==Decimal('.20') and values['cash']==Decimal('29.80')
    assert view.rows[1].values['日期']==('2026-09-30' if reverse else '2026-09-01')


@pytest.mark.parametrize('case',['no_total','two_totals','removed_last_day_with_stale_total','removed_middle_day_with_stale_total','duplicate_date','wrong_month','wrong_total'])
def test_corrupt_reports_are_rejected(case):
    rows=records()
    if case=='no_total':rows=rows[1:]
    if case=='two_totals':rows.append(deepcopy(rows[0]))
    if case=='removed_last_day_with_stale_total':rows.pop()
    if case=='removed_middle_day_with_stale_total':rows.pop(3)
    if case=='duplicate_date':rows.append(deepcopy(rows[1]))
    if case=='wrong_month':rows[1]['日期']='2026-08-31'
    if case=='wrong_total':rows[0]['非赠款消耗(元)']=999
    if case=='invalid_balance':rows[-1]['非赠款余额(元)']='-'
    with pytest.raises(DataError):statement([table(rows)],'2026-09')


@pytest.mark.parametrize('month,days',[('2024-02',29),('2025-02',28),('2026-12',31)])
def test_month_end_calendar_boundaries(month,days):
    _,values=statement([table(records(month,days))],month)
    assert values['date']==f'{month}-{days}'


def test_report_keeps_source_data_screenshot_and_exact_template_formulas(tmp_path):
    source=tmp_path/'日结报表.xlsx'
    wb=openpyxl.Workbook();ws=wb.active;ws.title='日结报表';ws.append(COLUMNS)
    for r in records():ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
    wb.save(source);wb.close()
    shot=tmp_path/'明细.png';Image.new('RGB',(1920,1080),'#345678').save(shot)
    manifest=tmp_path/'collection.json'
    manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'qianchuan','account':'千川账户'}],
                                   'screenshots':[{'path':str(shot),'sheet':'千川账户日结报表'}],
                                   'balances':[{'account':'千川账户','period':'2026-09','kind':'monthly','opening':'3','source':{'manual_verified':True}}]}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='测试店',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report']);template=openpyxl.load_workbook(TEMPLATE)
    assert book.sheetnames==template.sheetnames
    ws=book['千川账户']
    assert [ws.cell(1,c).value for c in range(1,15)]==COLUMNS
    assert ws['A2'].value=='总计' and ws['C2'].value==60
    assert ws['A32'].value=='2026-09-30' and ws['N32'].value==29.8
    assert ws['Q4'].value==.2 and ws['R4'].value==29.8
    assert len(ws._images)==1
    assert ws._images[0]._data()==shot.read_bytes()
    assert ws._images[0].anchor._from.col>=15 and ws._images[0].anchor._from.row>=7
    for coord,formula in TOTAL_FORMULAS.items():
        got=book['核对汇总'][coord].value
        assert got.text==formula and got.ref==coord
    for coord in ['L6','N6','O6','P6']:assert book['核对汇总'][coord].value==template['核对汇总'][coord].value
    book.close();template.close()
    cached=openpyxl.load_workbook(result['report'],data_only=True)
    assert cached['核对汇总']['D44'].value==-60
    assert cached['核对汇总']['D45'].value==87
    assert cached['核对汇总']['M6'].value==3
    assert cached['核对汇总']['N6'].value==27
    assert cached['核对汇总']['O6'].value==30
    assert cached['核对汇总']['P6'].value==0
    cached.close()
    root=Path(result['run_dir']);result_json=json.loads((root/'result.json').read_text())
    assert Decimal(next(c for c in result_json['checks'] if c['项目']=='千川账户余额')['差异'])==0
    schemas=json.loads((root/'tables.json').read_text())['tables']
    assert next(s for s in schemas if s['table_id']=='qianchuan_monthly')['excel_sheet']=='千川账户'


def test_missing_source_never_reuses_example_balances_or_totals(tmp_path):
    result=build([],month='2026-09',shop_name='缺少数据',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report'],data_only=True)
    assert book['千川账户']['C2'].value is None
    assert book['千川账户']['Q4'].value=='#N/A' and book['千川账户']['R4'].value=='#N/A'
    assert book['核对汇总']['D44'].value=='#N/A' and book['核对汇总']['P6'].value=='#N/A'
    assert not book['千川账户']._images
    book.close()
    issues=json.loads((Path(result['run_dir'])/'result.json').read_text())['issues']
    assert any(i['code']=='qianchuan_statement_invalid' for i in issues)


def test_truncated_download_with_stale_total_cannot_supply_month_end(tmp_path):
    source=tmp_path/'incomplete.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.append(COLUMNS)
    for r in records()[:-1]:ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
    wb.save(source);wb.close()
    result=build([source],month='2026-09',shop_name='缺少月末',output=tmp_path/'runs',template=TEMPLATE)
    book=openpyxl.load_workbook(result['report'],data_only=True)
    assert book['千川账户']['C2'].value==60
    assert book['千川账户']['Q4'].value=='#N/A'
    assert book['千川账户']['R4'].value=='#N/A'
    assert book['千川账户']['A31'].value is not None
    book.close()
    issues=json.loads((Path(result['run_dir'])/'result.json').read_text())['issues']
    assert any(i['code']=='qianchuan_statement_invalid' and '总计与每日明细不符' in i['message'] for i in issues)


def sparse_records(days):
    daily=[r for r in records()[1:] if r['日期'].day in days]
    total={'日期':'总计',**{k:sum((r[k] for r in daily),Decimal(0)) for k in COLUMNS[1:11]},
           **{k:'-' for k in COLUMNS[11:]}}
    return [total,*daily]


@pytest.mark.parametrize('days',[[10],[1,7,30],[20,23,29],[]])
def test_sparse_or_empty_statement_uses_actual_records_without_filling_dates(days):
    rows=sparse_records(days)
    view,result=statement([table(rows)],'2026-09',verified_counts={None:len(days)})
    assert len(view.rows)==len(days)+1
    assert result['consumption']==Decimal(-2*len(days))
    assert result['recharge']==Decimal('2.9')*len(days)
    assert result['cash']==(Decimal('29.8') if 30 in days else None)
    assert result['grant']==(Decimal('.2') if 30 in days else None)


def test_header_only_statement_requires_official_empty_query():
    with pytest.raises(DataError,match='只有表头'):statement([table([])],'2026-09')
    view,result=statement([table([])],'2026-09',verified_counts={None:0})
    assert len(view.rows)==1 and result['consumption']==0 and result['recharge']==0
    assert result['data_status']=='no_activity' and result['cash'] is None


def test_zero_total_alone_is_not_proof_of_no_activity():
    with pytest.raises(DataError,match='只有总计'):
        statement([table(sparse_records([]))],'2026-09')


def test_multiple_sparse_accounts_do_not_sum_unknown_balances_as_zero():
    first=table(sparse_records([7,30]));second=table(sparse_records([10]))
    for row in first.rows:row.source['entity_id']='one'
    for row in second.rows:row.source['entity_id']='two'
    view,result=statement([first,second],'2026-09')
    assert len(view.rows)==4 and result['consumption']==-6
    assert result['grant'] is None and result['cash'] is None
    assert all(row.values['非赠款余额(元)']=='-' for row in view.rows)


def test_sparse_statement_exports_totals_and_keeps_missing_month_end_normal(tmp_path):
    source=tmp_path/'sparse.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.append(COLUMNS)
    for r in sparse_records([10,20]):ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
    wb.save(source);wb.close()
    shot=tmp_path/'detail.png';Image.new('RGB',(2560,1080)).save(shot)
    manifest=tmp_path/'collection.json'
    from douyin_settle_recon.collect import KINDS
    receipts=[{'role':k,'status':'no_activity'} for k in KINDS if k!='qianchuan']
    receipts.append(q_receipt(total=2))
    manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'qianchuan','entity_id':'123'}],
                                   'collection':receipts,'screenshots':[{'path':str(shot),'sheet':'千川账户日结报表'}]}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='稀疏日期',output=tmp_path/'runs',template=TEMPLATE)
    report=json.loads((Path(result['run_dir'])/'result.json').read_text())
    assert report['issues']==[]
    check=next(c for c in report['checks'] if c['项目']=='千川账户余额')
    assert check['状态']=='无月末余额记录' and check['实际值'] is None
    book=openpyxl.load_workbook(result['report'],data_only=True)
    assert book['千川账户']['A2'].value=='总计' and book['千川账户']['A5'].value is None
    assert book['千川账户']['Q4'].value=='#N/A'
    assert book['千川账户']['Q5'].value=='月末无余额记录'
    assert book['核对汇总']['D44'].value==-4 and book['核对汇总']['D45'].value==5.8
    assert len(book['千川账户']._images)==1
    book.close()


def test_multiple_accounts_are_summed_once_by_day_and_account_identity():
    first=table(records());second=table(records())
    for row in first.rows:row.source.update(entity_id='one',entity_name='账户一')
    for row in second.rows:row.source.update(entity_id='two',entity_name='账户二')
    view,result=statement([first,second],'2026-09')
    assert len(view.rows)==31
    assert view.rows[0].values['日期']=='总计'
    assert view.rows[0].values['非赠款消耗(元)']==120
    assert result['grant']==Decimal('.40') and result['cash']==Decimal('59.60')
    assert result['consumption']==-120 and result['recharge']==174
    with pytest.raises(DataError,match='不同'):statement([first,first],'2026-09')


def test_multiple_accounts_keep_both_raw_statements_and_screenshots(tmp_path):
    sources=[];shots=[]
    for ident in ['one','two']:
        source=tmp_path/f'{ident}.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.append(COLUMNS)
        for r in records():ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
        wb.save(source);wb.close()
        sources.append({'path':str(source),'role':'qianchuan','entity_id':ident,'entity_name':ident})
        shot=tmp_path/f'{ident}.png';Image.new('RGB',(2560,1080)).save(shot)
        shots.append({'path':str(shot),'sheet':'千川账户日结报表-'+ident,'entity_id':ident,'entity_name':ident})
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'inputs':sources,'screenshots':shots}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='多账户',output=tmp_path/'runs',template=TEMPLATE)
    wb=openpyxl.load_workbook(result['report'],data_only=True);ws=wb['千川账户']
    assert ws['C2'].value==120 and ws['Q4'].value==.4 and ws['R4'].value==59.6
    assert len(ws._images)==2
    assert ws._images[1].anchor._from.row>ws._images[0].anchor.to.row
    assert sum(row[0]=='总计' for row in ws.values)==1
    assert sum(row[1]=='总计' for row in ws.values)==2
    assert sum(isinstance(row[0],str) and row[0].startswith('广告账户：') for row in ws.values)==2
    assert wb['核对汇总']['D44'].value==-120 and wb['核对汇总']['D45'].value==174
    wb.close()


def test_empty_account_capture_is_not_embedded_and_detail_capture_moves_up(tmp_path):
    source=tmp_path/'daily.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.append(COLUMNS)
    for r in records():ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
    wb.save(source);wb.close()
    shots=[]
    for ident,total,color in [('empty',0,'#FFFFFF'),('detail',30,'#336699')]:
        path=tmp_path/f'{ident}.png';Image.new('RGB',(2560,1080),color).save(path)
        shots.append({'path':str(path),'sheet':'千川账户日结报表-'+ident,'advertiser_id':ident,'query':{'total':total}})
    manifest=tmp_path/'input.json';manifest.write_text(json.dumps({'inputs':[{'path':str(source),'role':'qianchuan'}],'screenshots':shots}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='截图筛选',output=tmp_path/'runs',template=TEMPLATE)
    wb=openpyxl.load_workbook(result['report'],data_only=True);ws=wb['千川账户']
    assert len(ws._images)==1 and ws._images[0]._data()==(tmp_path/'detail.png').read_bytes()
    assert ws._images[0].anchor._from.row==7
    assert ws['C2'].value==60 and ws['Q4'].value==.2 and ws['R4'].value==29.8
    assert wb['核对汇总']['D44'].value==-60
    wb.close()
    handoff=json.loads((Path(result['run_dir'])/'manifest.json').read_text())
    empty=next(s for s in handoff['screenshots'] if s.get('advertiser_id')=='empty')
    assert empty['excel_sheet'] is None and empty['omitted_reason']=='empty_qianchuan_statement'
    assert (Path(result['run_dir'])/empty['path']).is_file()


def test_permission_denied_keeps_other_data_and_formulas_but_no_qianchuan_values(tmp_path):
    import csv
    from test_local_pipeline import fund,FUND_COLUMNS
    ledger=tmp_path/'fund.csv'
    row=fund(动账时间='2026-09-10 12:00:00')
    with ledger.open('w',encoding='utf-8-sig',newline='') as f:
        writer=csv.writer(f);writer.writerow(FUND_COLUMNS);writer.writerow([row.get(k) for k in FUND_COLUMNS])
    qfile=tmp_path/'partial-q.xlsx';wb=openpyxl.Workbook();ws=wb.active;ws.append(COLUMNS)
    for r in sparse_records([10]):ws.append([float(r[k]) if isinstance(r[k],Decimal) else r[k] for k in COLUMNS])
    wb.save(qfile);wb.close()
    shot=tmp_path/'partial.png';Image.new('RGB',(128,128)).save(shot)
    manifest=tmp_path/'collection.json'
    manifest.write_text(json.dumps({'inputs':[{'path':str(ledger),'role':'fund'},{'path':str(qfile),'role':'qianchuan'}],
        'collection':[{'role':'fund','status':'downloaded'},{'role':'qianchuan','status':'downloaded'},
                      {'role':'qianchuan','status':'permission_denied','error':'账户权限不足'}],
        'screenshots':[{'path':str(shot),'sheet':'千川账户日结报表'}]}))
    result=build([],manifest_path=manifest,month='2026-09',shop_name='权限留空测试',output=tmp_path/'runs',template=TEMPLATE)
    cached=openpyxl.load_workbook(result['report'],data_only=True)
    assert cached['资金账单'].max_row==2
    assert cached['千川账户']['A2'].value is None and not cached['千川账户']._images
    for coord in ['Q4','R4']:assert cached['千川账户'][coord].value is None
    for coord in ['D44','D45','M6','N6','O6','P6']:assert cached['核对汇总'][coord].value is None,coord
    cached.close()
    actual=openpyxl.load_workbook(result['report'],data_only=False);template=openpyxl.load_workbook(TEMPLATE,data_only=False)
    for row in template['核对汇总']:
        for cell in row:
            if cell.data_type=='f':
                got=actual['核对汇总'][cell.coordinate].value
                assert getattr(got,'text',got)==getattr(cell.value,'text',cell.value),cell.coordinate
    actual.close();template.close()
    payload=json.loads((Path(result['run_dir'])/'result.json').read_text())
    assert payload['context']['conclusion']=='待确认'
    assert any(i['code']=='qianchuan_permission_denied' for i in payload['issues'])
    assert not any(i['code']=='qianchuan_statement_invalid' for i in payload['issues'])


def test_missing_daily_balances_preserve_valid_expense_totals():
    rows=records();rows[-1]['赠款余额(元)']=None;rows[-1]['非赠款余额(元)']='-'
    _,result=statement([table(rows)],'2026-09')
    assert result['consumption']==-60 and result['recharge']==87
    assert result['grant'] is None and result['cash'] is None
