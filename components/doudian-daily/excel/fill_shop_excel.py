#!/usr/bin/env python3
"""把工作台采集模块 shopmetrics:metrics 的结果回填到抖店采集表。
用法: python3 fill_shop_excel.py <metrics_result.json> [输入xlsx] [输出xlsx]
"""
import json, sys, os
from datetime import datetime
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else '/tmp/metrics_result.json'
IN_XLSX = sys.argv[2] if len(sys.argv) > 2 else os.path.expanduser('~/Downloads/抖店Playwright采集字段导出(1).xlsx')
OUT_XLSX = sys.argv[3] if len(sys.argv) > 3 else os.path.expanduser('~/Downloads/抖店Playwright采集字段导出(1)-已填充.xlsx')

MISSING = '❎'


def load_rows(path):
    g = json.load(open(path, encoding='utf-8'))
    res = (g.get('result') or g).get('results') or []
    return [r for r in res if r.get('status') == 'ok']


def pct(v):
    return None if v is None else round(float(v) * 100, 4)


def build_map(r):
    """采集字段 -> 表列值"""
    sales = r.get('sales')
    ad = r.get('ad_cost')
    exp = r.get('exposure_uv')
    pay = r.get('pay_amt')
    m = {
        '平台': '抖店',
        '订单数': r.get('orders'),
        '客单价(元)': r.get('aov'),
        '店铺商品数': r.get('product_onsale'),
        '店铺总体验分': r.get('exp_total'),
        '商品体验分': r.get('exp_goods'),
        '服务体验分': r.get('exp_service'),
        '物流体验分': r.get('exp_logistics'),
        # 体验分项（店铺表现 = 平台原始值）
        '商品品质退货率（店铺表现）': r.get('m_quality_return'),
        '商品品质退货率（得分）': r.get('s_quality_return'),
        '运单配送时效达成率（店铺表现）': r.get('m_delivery_ontime'),
        '运单配送时效达成率（得分）': r.get('s_delivery_ontime'),
        '揽收时长平均得分': r.get('m_collect_duration'),
        '发货物流品退率（店铺表现）': r.get('m_logistics_return'),
        '发货物流品退率（得分）': r.get('s_logistics_return'),
        '飞鸽平均响应时长（店铺表现）': r.get('m_feige_response'),
        '飞鸽平均响应时长（得分）': r.get('s_feige_response'),
        '售后平均审核时长（店铺表现）': r.get('m_aftersale_audit'),
        '售后平均审核时长（得分）': r.get('s_aftersale_audit'),
        '飞鸽会话不满意率（店铺表现）': pct(r.get('m_feige_dissatisfy')),
        '飞鸽会话不满意率（得分）': r.get('s_feige_dissatisfy'),
        '平台求助率（店铺表现）': pct(r.get('m_platform_help')),
        '平台求助率（得分）': r.get('s_platform_help'),
        '商品综合评分（店铺表现）': r.get('m_goods_rating'),
        '商品综合评分（得分）': r.get('s_goods_rating'),
        '影响消费者体验（违规扣分项）': r.get('m_affect_experience'),
        '虚假交易刷体验分（违规扣分项）': r.get('m_fake_trade'),
        # 评价
        '好评率': r.get('comment_good_rate'),
        '差评率': r.get('comment_bad_rate'),
        '差评数': r.get('comment_negative'),
        # 流量 / 转化
        '曝光量': exp,
        '点击率%': r.get('ctr'),
        '点击转化率': r.get('cvr'),
        # 费用
        '推广费用(元)': ad,
        # 订单口径
        '延迟发货件数': r.get('order_ship_delay'),
        '未发订单数(订单口径)': r.get('order_stock_up'),
        '店铺未发货件数': r.get('order_stock_up'),
        '上新数': r.get('listed_today') if r.get('listed_today') is not None else r.get('listed_7d'),
        '下架数': r.get('offline_today') if r.get('offline_today') is not None else r.get('offline_7d'),
        '店铺缺货件数': r.get('stock_sold_out'),
        '店铺急缺件数': r.get('stock_advised_increase'),
        '缺货预警': r.get('stock_alarm'),
        '未发件数(订单口径)': r.get('order_stock_up'),
        # 日期
        '日期': datetime.now().strftime('%Y-%m-%d'),
        '同步时间': r.get('captured_at'),
        '更新时间': r.get('captured_at'),
    }
    # 加工指标
    if sales is not None and ad:
        m['推广ROI'] = round(sales / ad, 4)
    if pay is not None and exp:
        m['千次曝光用户支付金额'] = round(pay / exp * 1000, 4)
    return m


def main():
    rows = load_rows(SRC)
    if not rows:
        print('没有成功采集的店铺，未生成文件')
        return
    wb = openpyxl.load_workbook(IN_XLSX)
    ws = wb['采集结果']
    hdr = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(hdr)}

    for row_no, r in enumerate(rows, start=2):
        while ws.max_row < row_no:
            ws.append([None] * len(hdr))
        vals = build_map(r)
        vals['店铺'] = r.get('name')
        filled = 0
        for h, i in idx.items():
            v = vals.get(h)
            if v is None or v == '':
                continue
            ws.cell(row=row_no, column=i + 1).value = v
            filled += 1
        print(f"  {r.get('name')}: 填 {filled} 列")
    wb.save(OUT_XLSX)
    print('已保存:', OUT_XLSX)


if __name__ == '__main__':
    main()
